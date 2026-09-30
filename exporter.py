"""
Exporter Module
===============
Exports student records into structured CSV, formatted Excel (.xlsx),
and Course Catalog JSON metadata files.
"""

import csv
import json
import os
import re
from typing import Any, Dict, List, Optional, Tuple

try:
    from .models import CourseCatalog, ExamMeta, StudentRecord
except Exception:
    try:
        from data_extractor_AI.models import CourseCatalog, ExamMeta, StudentRecord
    except Exception:
        from models import CourseCatalog, ExamMeta, StudentRecord

try:
    from core.excel_report_generator import AcademicExcelReportGenerator, get_academic_filename
except Exception:
    try:
        from .core.excel_report_generator import AcademicExcelReportGenerator, get_academic_filename
    except Exception:
        from data_extractor_AI.core.excel_report_generator import AcademicExcelReportGenerator, get_academic_filename


def sanitize_folder_name(name: str, max_len: int = 50) -> str:
    """Sanitize filename/stem for directory naming."""
    clean = re.sub(r'[\\/*?:"<>|]', "", name)
    clean = clean.strip().replace(" ", "_")
    if len(clean) > max_len:
        clean = clean[:max_len].rstrip("_.-")
    return clean or "ai_extracted_results"


def write_csv_file(data: List[Dict[str, Any]], fieldnames: List[str], output_path: str) -> None:
    """Write list of row dictionaries to UTF-8 CSV with Windows long-path safety."""
    abs_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(abs_path), exist_ok=True)
    win_path = abs_path
    if os.name == "nt" and not abs_path.startswith("\\\\?\\") and len(abs_path) >= 240:
        win_path = "\\\\?\\" + abs_path
    with open(win_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)


def write_excel_file(
    data: List[Dict[str, Any]],
    fieldnames: List[str],
    catalog: Optional[CourseCatalog],
    meta: ExamMeta,
    output_path: str
) -> None:
    """Write formatted Excel workbook with Extracted Data and Course Catalog sheets."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        return

    wb = openpyxl.Workbook()

    # --- Sheet 1: Extracted Data ---
    ws = wb.active
    ws.title = "Extracted Data"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.append(fieldnames)
    for col_idx in range(1, len(fieldnames) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for row_dict in data:
        row_vals = [row_dict.get(col, "") for col in fieldnames]
        ws.append(row_vals)

    ws.freeze_panes = "A2"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    # --- Sheet 2: Course Catalog Metadata ---
    ws_cat = wb.create_sheet("Course Catalog")
    cat_headers = ["Course Code", "Official Course Name", "Component Heads", "Max Marks", "Credits"]
    ws_cat.append(cat_headers)

    cat_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    for col_idx in range(1, len(cat_headers) + 1):
        cell = ws_cat.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = cat_header_fill
        cell.alignment = center_align
        cell.border = thin_border

    courses_list = catalog.courses if catalog else (meta.catalog.courses if meta and meta.catalog else [])
    for c in courses_list:
        heads_str = ", ".join(c.heads) if c.heads else "N/A"
        max_m_str = ", ".join(f"{h}:{m}" for h, m in c.max_marks.items()) if c.max_marks else "N/A"
        ws_cat.append([c.code, c.name, heads_str, max_m_str, c.credits or ""])

    for col in ws_cat.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws_cat.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 60)

    abs_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(abs_path), exist_ok=True)
    wb.save(abs_path)


def export_ai_extracted_data(
    meta: ExamMeta,
    students: List[StudentRecord],
    input_pdf_path: str,
    base_output_dir: str,
    archetype: str = "tabulation",
    checksum_data: Optional[Dict[str, Any]] = None,
    flat_rows: Optional[List[Dict[str, Any]]] = None
) -> Tuple[str, str]:
    """
    Export extracted StudentRecord objects to CSV, Excel, and JSON catalog.
    Returns (designated_folder_path, output_csv_path).
    """
    pdf_filename = os.path.basename(input_pdf_path)
    pdf_stem = os.path.splitext(pdf_filename)[0]
    folder_name = sanitize_folder_name(pdf_stem) or "ai_extracted_results"

    designated_dir = os.path.join(base_output_dir, folder_name)
    os.makedirs(designated_dir, exist_ok=True)

    csv_path = os.path.join(designated_dir, f"{folder_name}_extracted_data.csv")
    excel_path = os.path.join(designated_dir, f"{folder_name}_extracted_data.xlsx")
    json_path = os.path.join(designated_dir, f"{folder_name}_catalog.json")

    # Flatten students to dicts if not provided
    if flat_rows is None:
        flat_rows = [s.to_flat_dict(meta.catalog) for s in students]

    # Construct canonical field names
    fieldnames: List[str] = []
    seen_fields = set()

    def add_field(f: str):
        if f not in seen_fields:
            seen_fields.add(f)
            fieldnames.append(f)

    # 1. Base student identity & academic status
    base_headers = [
        "Seat_No", "Student_Name", "Gender", "Candidate_Type", "PRN",
        "Semester", "Result_Status", "Remark", "Total_Marks",
        "Credits_Earned", "Total_CG", "SGPI"
    ]
    # Check if ERN, Mother_Name, ABC_ID exist in data
    if any("Mother_Name" in r and r["Mother_Name"] for r in flat_rows):
        base_headers.insert(5, "Mother_Name")
    if any("ERN" in r and r["ERN"] for r in flat_rows):
        base_headers.insert(5, "ERN")
    if any("ABC_ID" in r and r["ABC_ID"] for r in flat_rows):
        base_headers.insert(5, "ABC_ID")
    if any("CGPI" in r and r["CGPI"] for r in flat_rows):
        base_headers.append("CGPI")

    for bh in base_headers:
        add_field(bh)

    # 2. Courses in exact catalog order with only authorized heads
    if meta and meta.catalog and meta.catalog.courses:
        for c in meta.catalog.courses:
            # Only add heads that physically exist for this course (no default to TH, IA, TOT)
            actual_heads = [h for h in c.heads if h and str(h).strip()] if c.heads else []
            for h in actual_heads:
                add_field(f"{c.code}_{h}")
            
            add_field(f"{c.code}_Grade")
            # Only add GP, Credits, CG if credits > 0 or if present in candidate records
            has_credits = getattr(c, "credits", 0) and float(c.credits) > 0
            has_gp_data = any(str(r.get(f"{c.code}_GP", "")).strip() not in ("", "None", "0") for r in flat_rows)
            has_cg_data = any(str(r.get(f"{c.code}_CG", "")).strip() not in ("", "None", "0", "0.0") for r in flat_rows)
            
            if has_credits or has_gp_data:
                add_field(f"{c.code}_GP")
            if has_credits or any(str(r.get(f"{c.code}_Credits", "")).strip() not in ("", "None", "0") for r in flat_rows):
                add_field(f"{c.code}_Credits")
            if has_credits or has_cg_data:
                add_field(f"{c.code}_CG")
            add_field(f"{c.code}_Subject")

    # 3. Add any historical or remaining custom fields from flat rows (excluding hallucinated heads)
    for r in flat_rows:
        for k in r.keys():
            if k not in seen_fields:
                # Do not add phantom course head columns if catalog is defined
                if meta and meta.catalog and meta.catalog.courses and "_" in k:
                    prefix, suffix = k.split("_", 1)
                    c_match = meta.catalog.get_by_code(prefix)
                    if c_match:
                        valid_heads = set(c_match.heads) if c_match.heads else set()
                        if suffix in ("TH", "IA", "TW", "PR", "TOT", "OR", "ESE", "PO", "PRA") and suffix not in valid_heads:
                            continue
                add_field(k)

    # 4. Global Empty-Column Pruner:
    # Completely drop any column that is 100% empty across all student rows
    # (keeps essential identification columns: Seat_No, Student_Name, Result_Status, Semester)
    protected_core = {"Seat_No", "Student_Name", "Result_Status", "Semester"}
    pruned_fieldnames: List[str] = []
    for f in fieldnames:
        if f in protected_core:
            pruned_fieldnames.append(f)
        else:
            has_data = any(
                r.get(f) is not None
                and str(r.get(f)).strip() != ""
                and str(r.get(f)).strip() != "None"
                for r in flat_rows
            )
            if has_data:
                pruned_fieldnames.append(f)
    fieldnames = pruned_fieldnames

    # 1. Export CSV
    write_csv_file(flat_rows, fieldnames, csv_path)

    # 2. Export Excel
    write_excel_file(
        data=flat_rows,
        fieldnames=fieldnames,
        catalog=meta.catalog,
        meta=meta,
        output_path=excel_path
    )

    # 3. Export JSON Metadata Catalog
    cat_dict = meta.catalog.to_dict() if meta.catalog else []
    payload = {
        "exam_title": meta.exam_title,
        "semester": meta.semester,
        "college_name": meta.college_name,
        "result_date": meta.result_date,
        "archetype": meta.archetype or archetype,
        "total_students": len(students),
        "courses": cat_dict,
        "checksum_verification": checksum_data or {}
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)

    # 4. Export Academic Analysis 6-Sheet Workbook (Target Google Sheet format)
    try:
        meta_dict = {
            "academic_year": getattr(meta, "academic_year", ""),
            "semester": getattr(meta, "semester", ""),
            "exam_title": getattr(meta, "exam_title", ""),
            "college_name": getattr(meta, "college_name", "")
        }
        academic_filename = get_academic_filename(folder_name, meta_dict)
        academic_path = os.path.join(designated_dir, academic_filename)
        generator = AcademicExcelReportGenerator(
            students=flat_rows,
            catalog={"courses": cat_dict},
            meta=meta_dict,
            folder_name=folder_name,
            fieldnames=fieldnames
        )
        generator.generate_workbook(academic_path)
    except Exception as e:
        print(f"Warning: Academic analysis excel generation skipped: {e}")

    return designated_dir, csv_path
