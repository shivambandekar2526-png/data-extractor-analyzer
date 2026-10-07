"""
Unit and Integration Tests for Phase 5: Semester-wise Academic Analysis
======================================================================
Verifies:
1. Exact 3-tier dropdown hierarchy (Department -> Batch -> Semester Excel Sheet).
2. Deterministic calculation of Appeared, Passed, Failed, Pass %, Fail %.
3. Strict definition of Semester Success Rate = (Passed / Appeared) * 100.
4. Institutional Classifications (Distinction, First Class, Second Class, Pass Class).
5. SGPI Distribution calibrated across actual semester SGPI values.
6. Top 5 Semester Toppers ranked strictly by SGPI.
7. Word (.docx) document generation with 1-page institutional layout & tables.
8. API endpoints:
   - GET /analysis/semester
   - GET /api/semester-analysis/options
   - POST /api/semester-analysis/analyze
   - POST /api/semester-analysis/export-docx
"""

import json
import os
import sys
import unittest
import docx

APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from app import app, resolve_subject_dataset
from core.semester_analyzer import (
    clean_student_name,
    clean_seat_no,
    clean_sgpi_value,
    resolve_student_pass_status,
    classify_student_semester,
    compute_sgpi_distribution,
    analyze_semester_performance
)
from core.semester_docx_exporter import generate_semester_analysis_docx


class TestSemesterAcademicAnalysis(unittest.TestCase):

    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def test_sgpi_and_status_cleaning(self):
        """Verifies SGPI parser and result status determination."""
        # Clean SGPI
        self.assertEqual(clean_sgpi_value("9.50"), 9.50)
        self.assertEqual(clean_sgpi_value(8.75), 8.75)
        self.assertEqual(clean_sgpi_value("0.0"), 0.0)
        self.assertIsNone(clean_sgpi_value("--"))
        self.assertIsNone(clean_sgpi_value("ABS"))
        self.assertIsNone(clean_sgpi_value(None))
        self.assertIsNone(clean_sgpi_value("nan"))

        # Clean Name
        self.assertEqual(clean_student_name("/PATIL SNEHA"), "PATIL SNEHA")
        self.assertEqual(clean_student_name("  ROHIT   SHARMA  "), "ROHIT SHARMA")

        # Pass / Fail status resolution
        status, passed = resolve_student_pass_status("P", 8.5)
        self.assertEqual(status, "PASS")
        self.assertTrue(passed)

        status, passed = resolve_student_pass_status("PASS", 9.0)
        self.assertEqual(status, "PASS")
        self.assertTrue(passed)

        status, passed = resolve_student_pass_status("F", None)
        self.assertEqual(status, "FAIL")
        self.assertFalse(passed)

        status, passed = resolve_student_pass_status("FAIL", 0.0)
        self.assertEqual(status, "FAIL")
        self.assertFalse(passed)

        status, passed = resolve_student_pass_status("ABS", None)
        self.assertEqual(status, "ABSENT")
        self.assertFalse(passed)

    def test_semester_success_rate_formula(self):
        """Verifies strict adherence to Success Rate = (Passed / Appeared) * 100."""
        dummy_records = [
            {"Seat_No": "1", "Student_Name": "A", "Result_Status": "P", "SGPI": 9.0},
            {"Seat_No": "2", "Student_Name": "B", "Result_Status": "P", "SGPI": 8.0},
            {"Seat_No": "3", "Student_Name": "C", "Result_Status": "P", "SGPI": 7.5},
            {"Seat_No": "4", "Student_Name": "D", "Result_Status": "F", "SGPI": None},
        ]
        res = analyze_semester_performance(dummy_records)
        self.assertTrue(res["success"])
        sm = res["semester_summary"]

        self.assertEqual(sm["appeared"], 4)
        self.assertEqual(sm["passed"], 3)
        self.assertEqual(sm["failed"], 1)

        # 3 / 4 * 100 = 75.0%
        self.assertEqual(sm["pass_rate"], 75.0)
        self.assertEqual(sm["fail_rate"], 25.0)
        self.assertEqual(sm["success_rate"], 75.0)

    def test_sgpi_distribution_bins(self):
        """Verifies SGPI distribution bins students accurately."""
        sgpis = [9.5, 9.1, 8.8, 8.2, 7.6, 6.4, 5.5, 4.2]
        dist = compute_sgpi_distribution(sgpis)

        # 9.00-10.00: 2
        # 8.00-8.99: 2
        # 7.00-7.99: 1
        # 6.00-6.99: 1
        # 5.00-5.99: 1
        # < 5.00: 1
        counts = dist["chart"]["counts"]
        self.assertEqual(counts, [2, 2, 1, 1, 1, 1])
        self.assertEqual(dist["total_scored_students"], 8)

    def test_topper_ranking_by_sgpi(self):
        """Verifies Top 5 toppers are strictly passing students ranked by SGPI descending."""
        records = [
            {"Seat_No": "101", "Student_Name": "STUDENT 1", "Result_Status": "P", "SGPI": 7.5},
            {"Seat_No": "102", "Student_Name": "STUDENT 2", "Result_Status": "P", "SGPI": 9.8},
            {"Seat_No": "103", "Student_Name": "STUDENT 3", "Result_Status": "F", "SGPI": 10.0}, # Failed should be excluded
            {"Seat_No": "104", "Student_Name": "STUDENT 4", "Result_Status": "P", "SGPI": 8.5},
            {"Seat_No": "105", "Student_Name": "STUDENT 5", "Result_Status": "P", "SGPI": 9.2},
            {"Seat_No": "106", "Student_Name": "STUDENT 6", "Result_Status": "P", "SGPI": 8.0},
            {"Seat_No": "107", "Student_Name": "STUDENT 7", "Result_Status": "P", "SGPI": 6.5},
        ]
        res = analyze_semester_performance(records)
        toppers = res["top_students"]

        self.assertEqual(len(toppers), 5)
        self.assertEqual(toppers[0]["student_name"], "STUDENT 2")
        self.assertEqual(toppers[0]["sgpi"], 9.8)
        self.assertEqual(toppers[1]["student_name"], "STUDENT 5")
        self.assertEqual(toppers[1]["sgpi"], 9.2)
        self.assertEqual(toppers[2]["student_name"], "STUDENT 4")
        self.assertEqual(toppers[2]["sgpi"], 8.5)
        self.assertEqual(toppers[3]["student_name"], "STUDENT 6")
        self.assertEqual(toppers[3]["sgpi"], 8.0)
        self.assertEqual(toppers[4]["student_name"], "STUDENT 1")
        self.assertEqual(toppers[4]["sgpi"], 7.5)

    def test_semester_analysis_real_dataset_sem8(self):
        """Verifies semester analysis on real 24-25 Semester 8 Excel sheet."""
        dataset_id = "excel_Artificial_Intelligence_and_Data_Science/24-25/Semester_8.xlsx"
        records, catalog, meta = resolve_subject_dataset(dataset_id)
        self.assertTrue(len(records) > 0, "Semester 8 records must exist")

        res = analyze_semester_performance(records, dataset_meta=meta)
        self.assertTrue(res["success"])
        sm = res["semester_summary"]

        self.assertEqual(sm["appeared"], 65)
        self.assertEqual(sm["passed"], 65)
        self.assertEqual(sm["failed"], 0)
        self.assertEqual(sm["pass_rate"], 100.0)
        self.assertEqual(sm["success_rate"], 100.0)
        self.assertTrue(sm["first_class_and_dist_count"] > 0)
        self.assertEqual(len(res["top_students"]), 5)

    def test_semester_docx_exporter(self):
        """Verifies institutional Word document structure for semester analysis."""
        dummy_analysis = {
            "header": {
                "department": "Artificial Intelligence and Data Science",
                "branch": "Artificial Intelligence and Data Science",
                "batch": "2024–2025",
                "academic_year": "2024–25",
                "semester": "Semester VIII",
                "dataset_name": "Semester_8.xlsx"
            },
            "semester_summary": {
                "appeared": 65,
                "passed": 60,
                "failed": 5,
                "pass_rate": 92.31,
                "fail_rate": 7.69,
                "success_rate": 92.31,
                "first_class_count": 35,
                "first_class_rate": 53.85,
                "distinction_count": 25,
                "distinction_rate": 38.46,
                "first_class_and_dist_count": 60,
                "first_class_and_dist_pct": 92.31,
                "second_class_count": 0,
                "pass_class_count": 0,
                "classification": {
                    "table": [
                        {"short": "Distinction", "classification": "First Class with Distinction", "students": 25, "percentage": 41.67},
                        {"short": "First Class", "classification": "First Class", "students": 35, "percentage": 58.33}
                    ]
                }
            },
            "sgpi_distribution": {
                "table": [
                    {"range": "9.00 – 10.00", "students": 15, "percentage": 25.0},
                    {"range": "8.00 – 8.99", "students": 30, "percentage": 50.0},
                    {"range": "7.00 – 7.99", "students": 15, "percentage": 25.0}
                ],
                "chart": {
                    "labels": ["9.00-10.00", "8.00-8.99", "7.00-7.99"],
                    "counts": [15, 30, 15]
                }
            },
            "top_students": [
                {
                    "rank": 1,
                    "student_name": "PATIL NIRAJ SANTOSH SUMAN",
                    "seat_or_prn": "12345",
                    "sgpi": 9.71,
                    "percentage": 88.5
                }
            ]
        }

        buf = generate_semester_analysis_docx(
            dummy_analysis,
            department_name="Artificial Intelligence and Data Science",
            batch_name="2024–2025",
            semester_name="Semester VIII"
        )
        self.assertIsNotNone(buf)
        doc = docx.Document(buf)
        self.assertGreaterEqual(len(doc.tables), 4)

        full_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        self.assertIn("TERNA ENGINEERING COLLEGE", full_text)
        self.assertIn("SEMESTER-WISE EXAMINATION PERFORMANCE", full_text)
        self.assertIn("Semester Success Rate", full_text)
        self.assertIn("PATIL NIRAJ SANTOSH SUMAN", full_text)

    def test_api_options_endpoint(self):
        """Verifies GET /api/semester-analysis/options returns 3-tier hierarchy."""
        res = self.app.get("/api/semester-analysis/options")
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["success"])
        self.assertIn("departments", data)
        self.assertGreater(len(data["departments"]), 0)

        # Department must contain batches which contain semester files
        first_dept = data["departments"][0]
        self.assertIn("batches", first_dept)
        if first_dept["batches"]:
            first_batch = first_dept["batches"][0]
            self.assertIn("semesters", first_batch)

    def test_api_analyze_endpoint(self):
        """Verifies POST /api/semester-analysis/analyze returns full analysis."""
        payload = {
            "dataset_id": "excel_Artificial_Intelligence_and_Data_Science/24-25/Semester_8.xlsx"
        }
        res = self.app.post("/api/semester-analysis/analyze", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["success"])
        self.assertIn("semester_summary", data)
        self.assertIn("sgpi_distribution", data)
        self.assertIn("top_students", data)
        self.assertEqual(data["semester_summary"]["appeared"], 65)
        self.assertEqual(data["semester_summary"]["success_rate"], 100.0)

    def test_api_export_docx_endpoint(self):
        """Verifies POST /api/semester-analysis/export-docx returns downloadable Word document."""
        payload = {
            "dataset_id": "excel_Artificial_Intelligence_and_Data_Science/24-25/Semester_8.xlsx"
        }
        res = self.app.post("/api/semester-analysis/export-docx", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res.headers["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        self.assertTrue(len(res.data) > 1000)

    def test_web_route_available(self):
        """Verifies GET /analysis/semester renders the application shell."""
        res = self.app.get("/analysis/semester")
        self.assertEqual(res.status_code, 200)
        html = res.get_data(as_text=True)
        self.assertIn("Semester-wise Analysis", html)
        self.assertIn("viewSemesterAnalysis", html)


if __name__ == "__main__":
    unittest.main()
