"""
Unit and Integration Tests for Phase 4: Subject-wise Academic Analysis
======================================================================
Verifies:
1. Dynamic subject derivation without hardcoding.
2. Deterministic calculation of appeared, passed, failed, pass %, fail %.
3. Institutional classification (Distinction, First Class, Second Class, Pass Class).
4. Adaptive marks distribution calibrated to maximum subject score.
5. Term Work and Oral/Practical detection and unavailable fallback notices.
6. Top 5 student ranking by total subject marks (not SGPI).
7. Data validation (0 marks vs Absent vs Missing vs Fail flags).
8. Full API integration across multiple real historical datasets.
"""

import json
import os
import sys
import unittest
import pandas as pd

APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from app import app
from core.academic_classifier import (
    AcademicClassificationRules,
    clean_mark_entry,
    compute_marks_distribution,
    extract_subject_definitions,
    analyze_subject_performance
)


class TestSubjectAcademicAnalysis(unittest.TestCase):

    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def test_clean_mark_entry_data_validation(self):
        """Verifies distinct handling of 0 marks, Absent, Missing, and Fail flags."""
        # Valid 0 mark (not absent, not missing)
        st, val, flag = clean_mark_entry(0)
        self.assertEqual(st, "VALID")
        self.assertEqual(val, 0.0)
        self.assertFalse(flag)

        st, val, flag = clean_mark_entry("0.0")
        self.assertEqual(st, "VALID")
        self.assertEqual(val, 0.0)

        # Absent students
        st, val, flag = clean_mark_entry("ABS")
        self.assertEqual(st, "ABSENT")
        self.assertIsNone(val)

        st, val, flag = clean_mark_entry("AB")
        self.assertEqual(st, "ABSENT")

        # Missing values
        st, val, flag = clean_mark_entry(None)
        self.assertEqual(st, "MISSING")

        st, val, flag = clean_mark_entry("--")
        self.assertEqual(st, "MISSING")

        st, val, flag = clean_mark_entry("")
        self.assertEqual(st, "MISSING")

        # Suffix F / fail flagged mark
        st, val, flag = clean_mark_entry("18F")
        self.assertEqual(st, "VALID")
        self.assertEqual(val, 18.0)
        self.assertTrue(flag)

        # Normal passing mark
        st, val, flag = clean_mark_entry(65.5)
        self.assertEqual(st, "VALID")
        self.assertEqual(val, 65.5)
        self.assertFalse(flag)

    def test_classification_rules(self):
        """Verifies standard classification rules without inventing thresholds."""
        self.assertEqual(AcademicClassificationRules.classify_percentage(85.0), "First Class with Distinction")
        self.assertEqual(AcademicClassificationRules.classify_percentage(75.0), "First Class with Distinction")
        self.assertEqual(AcademicClassificationRules.classify_percentage(70.0), "First Class")
        self.assertEqual(AcademicClassificationRules.classify_percentage(60.0), "First Class")
        self.assertEqual(AcademicClassificationRules.classify_percentage(55.0), "Second Class")
        self.assertEqual(AcademicClassificationRules.classify_percentage(50.0), "Second Class")
        self.assertEqual(AcademicClassificationRules.classify_percentage(45.0), "Pass Class")
        self.assertEqual(AcademicClassificationRules.classify_percentage(40.0), "Pass Class")
        self.assertEqual(AcademicClassificationRules.classify_percentage(35.0), "Fail")

    def test_dynamic_subject_derivation_2122_8th(self):
        """Verifies subjects are dynamically discovered and paired with labs from 2122-8th."""
        csv_path = os.path.join(APP_DIR, "output", "2122-8th", "2122-8th_extracted_data.csv")
        cat_path = os.path.join(APP_DIR, "output", "2122-8th", "2122-8th_catalog.json")

        self.assertTrue(os.path.isfile(csv_path), "2122-8th CSV must exist")
        df = pd.read_csv(csv_path)
        with open(cat_path, mode="r", encoding="utf-8") as f:
            cat = json.load(f)

        subjects = extract_subject_definitions(df, cat)
        codes = [s["code"] for s in subjects]

        # Theory courses are present
        self.assertIn("CSC801", codes)
        self.assertIn("CSDO8013", codes)
        self.assertIn("AIDSCSP801", codes)
        # Standalone lab CSL801 was merged into CSC801
        self.assertNotIn("CSL801", codes)

        csc801 = next(s for s in subjects if s["code"] == "CSC801")
        self.assertEqual(csc801["name"], "Advanced Artificial Intelligence")
        self.assertTrue(csc801["has_theory"])
        # Lab components (TW, PR) were merged into the theory course
        self.assertTrue(csc801["has_term_work"])
        self.assertTrue(csc801["has_practical_oral"])
        self.assertTrue(csc801.get("is_combined"))
        self.assertEqual(csc801.get("lab_code"), "CSL801")

    def test_deterministic_theory_metrics_csc801(self):
        """Verifies deterministic pass/fail calculation, first class count, and combined ranking for CSC801."""
        csv_path = os.path.join(APP_DIR, "output", "2122-8th", "2122-8th_extracted_data.csv")
        cat_path = os.path.join(APP_DIR, "output", "2122-8th", "2122-8th_catalog.json")

        df = pd.read_csv(csv_path)
        with open(cat_path, mode="r", encoding="utf-8") as f:
            cat = json.load(f)

        res = analyze_subject_performance(df.to_dict(orient="records"), cat, "CSC801")

        self.assertTrue(res["success"])
        self.assertEqual(res["header"]["subject_code"], "CSC801 / CSL801")
        self.assertEqual(res["header"]["subject_name"], "Advanced Artificial Intelligence")
        # Combined max marks: 100 (theory) + 50 (lab) = 150
        self.assertEqual(res["header"]["total_max_marks"], 150.0)

        # Theory stats
        th = res["theory"]
        self.assertTrue(th["available"])
        self.assertEqual(th["appeared"], 65)
        self.assertEqual(th["passed"], 65)
        self.assertEqual(th["failed"], 0)
        self.assertEqual(th["pass_rate"], 100.0)
        self.assertEqual(th["fail_rate"], 0.0)
        self.assertIn("first_class_and_dist_count", th)
        self.assertEqual(th["first_class_and_dist_count"], 47)
        self.assertAlmostEqual(th["first_class_and_dist_pct"], 72.31, places=1)

        # Pass / Fail table contains First Class & Distinction
        pf_table = th["pass_fail_table"]
        pf_statuses = [r["status"] for r in pf_table]
        self.assertIn("Passed", pf_statuses)
        self.assertIn("Failed", pf_statuses)
        self.assertIn("First Class & Distinction (≥60%)", pf_statuses)

        # Classification table contains Distinction, First Class, Second Class, Pass Class
        class_table = th["classification"]["table"]
        c_names = [c["short"] for c in class_table]
        self.assertIn("Distinction", c_names)
        self.assertIn("First Class", c_names)
        self.assertIn("Second Class", c_names)
        self.assertIn("Pass Class", c_names)

        # Lab components (TW, PR) are merged and available
        self.assertTrue(res["term_work"]["available"])
        self.assertEqual(res["term_work"]["appeared"], 65)
        self.assertTrue(res["practical_oral"]["available"])
        self.assertEqual(res["practical_oral"]["appeared"], 65)

        # Viva distribution section is removed (not in response)
        self.assertNotIn("oral", res)

        # Top 5 students ranked based on total combined marks (out of 150)
        top_5 = res["top_students"]
        self.assertEqual(len(top_5), 5)
        self.assertEqual(top_5[0]["rank"], 1)
        self.assertEqual(top_5[0]["max_marks"], 150.0)
        self.assertTrue(top_5[0]["marks"] >= top_5[1]["marks"])
        self.assertTrue(top_5[0]["marks"] > 100.0, "Top mark out of 150 should exceed theory 100")

    def test_combined_subject_and_lab_pairing(self):
        """Verifies pairing of Reinforcement Learning (CSDO8013) with its lab (AIDSCSDOL8013)."""
        csv_path = os.path.join(APP_DIR, "output", "2122-8th", "2122-8th_extracted_data.csv")
        cat_path = os.path.join(APP_DIR, "output", "2122-8th", "2122-8th_catalog.json")

        df = pd.read_csv(csv_path)
        with open(cat_path, mode="r", encoding="utf-8") as f:
            cat = json.load(f)

        res = analyze_subject_performance(df.to_dict(orient="records"), cat, "CSDO8013")

        self.assertTrue(res["success"])
        self.assertEqual(res["header"]["subject_code"], "CSDO8013 / AIDSCSDOL8013")
        self.assertEqual(res["header"]["subject_name"], "Reinforcement Learning")
        self.assertEqual(res["header"]["total_max_marks"], 150.0)

        # Both theory and lab are active in the combined analysis
        self.assertTrue(res["theory"]["available"])
        self.assertTrue(res["term_work"]["available"])
        self.assertTrue(res["practical_oral"]["available"])

        # Check toppers are computed on combined score out of 150
        top_1 = res["top_students"][0]
        self.assertEqual(top_1["student_name"], "PATIL NIRAJ SANTOSH SUMAN")
        self.assertEqual(top_1["marks"], 133.0)
        self.assertEqual(top_1["max_marks"], 150.0)
        self.assertAlmostEqual(top_1["percentage"], 88.67, places=1)

    def test_multi_dataset_fail_tracking_sem3(self):
        """Verifies real failure tracking and term work on 23-24 Semester 3 Excel dataset."""
        excel_path = os.path.join(APP_DIR, "departments_data", "excel", "Artificial_Intelligence_and_Data_Science", "23-24", "Semester_3.xlsx")
        self.assertTrue(os.path.isfile(excel_path))

        xl = pd.ExcelFile(excel_path)
        df = xl.parse("Extracted Data")
        cat_df = xl.parse("Course Catalog") if "Course Catalog" in xl.sheet_names else None
        courses = []
        if cat_df is not None:
            for _, r in cat_df.iterrows():
                courses.append({"code": r.get("Course Code"), "name": r.get("Official Course Name")})
        cat = {"courses": courses}

        res = analyze_subject_performance(df.to_dict(orient="records"), cat, "CSC301")

        self.assertTrue(res["success"])
        th = res["theory"]
        self.assertTrue(th["available"])
        self.assertEqual(th["appeared"], 86)
        self.assertEqual(th["passed"], 65)
        self.assertEqual(th["failed"], 21)
        self.assertAlmostEqual(th["pass_rate"], 75.58, places=1)
        self.assertAlmostEqual(th["fail_rate"], 24.42, places=1)

        # Term work is also directly present in CSC301
        self.assertTrue(res["term_work"]["available"])
        self.assertEqual(res["term_work"]["appeared"], 87)

    def test_api_options_endpoint(self):
        """Tests GET /api/subject-analysis/options endpoint."""
        res = self.app.get("/api/subject-analysis/options")
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["success"])
        self.assertIn("departments", data)
        self.assertIn("academic_years", data)
        self.assertIn("branches", data)
        self.assertIn("semesters", data)
        self.assertIn("datasets", data)
        self.assertTrue(len(data["datasets"]) > 0)
        self.assertTrue(len(data["departments"]) > 0)

    def test_department_excel_hierarchy_and_smart_subjects(self):
        """Verifies department -> batch -> semester sheet -> smart subjects selection flow."""
        # 1. Options returns department hierarchy
        res = self.app.get("/api/subject-analysis/options")
        data = res.get_json()
        depts = data["departments"]
        self.assertTrue(len(depts) > 0)

        # Find AIDS department
        aids_dept = next((d for d in depts if "Artificial" in d["name"]), None)
        self.assertIsNotNone(aids_dept)
        self.assertEqual(aids_dept["folder"], "Artificial_Intelligence_and_Data_Science")

        # Find 21-22 batch inside it
        batch_2122 = next((b for b in aids_dept["batches"] if b["folder"] == "21-22"), None)
        self.assertIsNotNone(batch_2122)

        # Find Semester_1.xlsx inside it
        sem1_sheet = next((s for s in batch_2122["semesters"] if s["file_name"] == "Semester_1.xlsx"), None)
        self.assertIsNotNone(sem1_sheet)
        self.assertTrue(sem1_sheet["dataset_id"].startswith("excel_"))

        # 2. Query dynamic subjects smartly from Excel sheet
        sub_res = self.app.post("/api/subject-analysis/subjects", json={"dataset_id": sem1_sheet["dataset_id"]})
        self.assertEqual(sub_res.status_code, 200)
        sub_data = sub_res.get_json()
        self.assertTrue(sub_data["success"])
        self.assertTrue(len(sub_data["subjects"]) > 0)

        sub_codes = [s["code"] for s in sub_data["subjects"]]
        # Must contain valid courses from _subject columns
        self.assertIn("58651", sub_codes)
        self.assertIn("FEC101", sub_codes)
        # Must NOT contain non-course noise like MAX
        self.assertNotIn("MAX", sub_codes)

        # Subject title must be smartly populated from the _subject column
        m1 = next(s for s in sub_data["subjects"] if s["code"] == "58651")
        self.assertEqual(m1["name"], "Engineering Mathematics-I")
        self.assertEqual(m1["display_name"], "58651 - Engineering Mathematics-I")

        # 3. Analyze chosen subject from Excel sheet
        an_res = self.app.post("/api/subject-analysis/analyze", json={
            "dataset_id": sem1_sheet["dataset_id"],
            "subject_code": "58651"
        })
        self.assertEqual(an_res.status_code, 200)
        an_data = an_res.get_json()
        self.assertTrue(an_data["success"])
        self.assertEqual(an_data["header"]["subject_code"], "58651")
        self.assertEqual(an_data["theory"]["appeared"], 67)
        self.assertEqual(an_data["theory"]["passed"], 66)
        self.assertAlmostEqual(an_data["theory"]["pass_rate"], 98.51, places=1)
        self.assertEqual(len(an_data["top_students"]), 5)

    def test_api_subjects_endpoint(self):
        """Tests POST /api/subject-analysis/subjects endpoint."""
        res = self.app.post("/api/subject-analysis/subjects", json={"dataset_id": "out_2122-8th"})
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["success"])
        self.assertTrue(len(data["subjects"]) > 0)
        codes = [s["code"] for s in data["subjects"]]
        self.assertIn("CSC801", codes)

    def test_api_analyze_endpoint(self):
        """Tests POST /api/subject-analysis/analyze endpoint."""
        payload = {
            "dataset_id": "out_2122-8th",
            "subject_code": "CSC801",
            "academic_year": "2024–25",
            "branch": "Computer Engineering",
            "semester": "VIII"
        }
        res = self.app.post("/api/subject-analysis/analyze", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["success"])
        self.assertEqual(data["header"]["subject_code"], "CSC801 / CSL801")
        self.assertEqual(data["theory"]["appeared"], 65)
        self.assertIn("first_class_and_dist_count", data["theory"])
        self.assertEqual(len(data["top_students"]), 5)

    def test_distinct_th_pr_or_detection_and_analysis(self):
        """Verifies TH, PR, OR column detection and analysis are independent and accurate."""
        from core.academic_classifier import find_column_for_head

        cols = ["CSC801 - TH", "CSC801_PR", "CSC801 (OR)", "CSC801_TW", "CSC801_Subject"]
        self.assertEqual(find_column_for_head(cols, "CSC801", "TH"), "CSC801 - TH")
        self.assertEqual(find_column_for_head(cols, "CSC801", "PR"), "CSC801_PR")
        self.assertEqual(find_column_for_head(cols, "CSC801", "OR"), "CSC801 (OR)")
        self.assertEqual(find_column_for_head(cols, "CSC801", "TW"), "CSC801_TW")

        records = [
            {"Seat_No": "S1", "Student_Name": "Alice", "CSC801 - TH": 60, "CSC801_PR": 22, "CSC801 (OR)": 24, "CSC801_Subject": "AI"},
            {"Seat_No": "S2", "Student_Name": "Bob", "CSC801 - TH": 30, "CSC801_PR": 18, "CSC801 (OR)": 8, "CSC801_Subject": "AI"},
            {"Seat_No": "S3", "Student_Name": "Charlie", "CSC801 - TH": "ABS", "CSC801_PR": "ABS", "CSC801 (OR)": 15, "CSC801_Subject": "AI"}
        ]
        res = analyze_subject_performance(records, None, "CSC801")
        self.assertTrue(res["success"])

        # Theory
        self.assertTrue(res["theory"]["available"])
        self.assertEqual(res["theory"]["appeared"], 2)
        self.assertEqual(res["theory"]["passed"], 1)
        self.assertEqual(res["theory"]["failed"], 1)

        # Practical / Lab component
        self.assertTrue(res["practical_oral"]["available"])
        self.assertEqual(res["practical_oral"]["appeared"], 2)
        self.assertEqual(res["practical_oral"]["passed"], 2)

        # First class & distinction in Pass/Fail table
        pf_table = res["theory"]["pass_fail_table"]
        pf_names = [r["status"] for r in pf_table]
        self.assertIn("First Class & Distinction (≥60%)", pf_names)

        # Oral/Viva distribution section removed
        self.assertNotIn("oral", res)

    def test_naming_anomalies_and_ghost_pruning(self):
        """Verifies spaced codes merge properly and trailing heads don't create ghost subjects."""
        cols = ["FEC201_TH", "FEC201_Subject", "FEC201 TW_Credits", "FEL10 1_TH_OR", "FEL10 1_Subject"]
        df = pd.DataFrame({
            "FEC201_TH": [45, 55],
            "FEC201_Subject": ["Maths II", "Maths II"],
            "FEC201 TW_Credits": [1, 1],
            "FEL10 1_TH_OR": [20, 22],
            "FEL10 1_Subject": ["Physics Lab", "Physics Lab"]
        })
        subjs = extract_subject_definitions(df)
        codes = [s["code"] for s in subjs]
        self.assertIn("FEC201", codes)
        self.assertIn("FEL101", codes)
        self.assertNotIn("FEC201TW", codes)

    def test_live_folder_watcher_endpoints(self):
        """Verifies GET /api/live-watcher/status and POST /api/live-watcher/check-now."""
        res = self.app.get("/api/live-watcher/status")
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["success"])
        self.assertIn("version", data)
        self.assertIn("watched_categories", data)

        res_post = self.app.post("/api/live-watcher/check-now")
        self.assertEqual(res_post.status_code, 200)
        post_data = res_post.get_json()
        self.assertTrue(post_data["success"])

    def test_docx_exporter_structure(self):
        """Verifies institutional Word document structure for Terna Engineering College."""
        from core.subject_docx_exporter import generate_subject_analysis_docx
        import docx

        dummy_analysis = {
            "header": {
                "subject_code": "CSC801 / CSL801",
                "theory_code": "CSC801",
                "lab_code": "CSL801",
                "subject_name": "Advanced Artificial Intelligence",
                "department": "Artificial Intelligence and Data Science",
                "batch": "2021–2022",
                "semester": "Semester VIII",
                "academic_year": "2024–25",
                "total_max_marks": 150.0
            },
            "theory": {
                "appeared": 65,
                "passed": 65,
                "failed": 0,
                "pass_rate": 100.0,
                "fail_rate": 0.0,
                "first_class_count": 47,
                "first_class_rate": 72.31,
                "distinction_count": 18,
                "distinction_rate": 27.69,
                "first_class_and_dist_count": 65,
                "first_class_and_dist_rate": 100.0
            },
            "term_work": {
                "available": True,
                "appeared": 65,
                "passed": 65,
                "failed": 0,
                "pass_rate": 100.0
            },
            "practical_oral": {
                "available": True,
                "appeared": 65,
                "passed": 65,
                "failed": 0,
                "pass_rate": 100.0
            },
            "top_students": [
                {
                    "rank": 1,
                    "student_name": "PATIL NIRAJ SANTOSH SUMAN",
                    "seat_or_prn": "12345",
                    "marks": 133.0,
                    "max_marks": 150.0,
                    "percentage": 88.67
                }
            ]
        }

        buf = generate_subject_analysis_docx(
            analysis_data=dummy_analysis,
            department_name="Department of Artificial Intelligence and Data Science",
            batch_name="2021–2022",
            semester_name="Semester VIII"
        )
        self.assertGreater(len(buf.getvalue()), 1000)

        # Inspect generated docx contents
        doc = docx.Document(buf)
        full_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])

        # Header asserts
        self.assertIn("TERNA ENGINEERING COLLEGE", full_text)
        self.assertIn("DEPARTMENT OF ARTIFICIAL INTELLIGENCE AND DATA SCIENCE", full_text)
        self.assertIn("2021–2022", full_text)
        self.assertIn("Semester VIII", full_text)
        self.assertIn("CSC801", full_text)

        # Content asserts
        self.assertIn("Total Students Appeared", full_text)
        self.assertIn("Students Passed", full_text)
        self.assertIn("Students Failed", full_text)
        self.assertIn("Students with First Class", full_text)
        self.assertIn("Students with Distinction", full_text)
        self.assertIn("PATIL NIRAJ SANTOSH SUMAN", full_text)
        self.assertIn("Head of Department", full_text)

    def test_export_docx_api_endpoint(self):
        """Verifies POST /api/subject-analysis/export-docx returns downloadable Word document."""
        payload = {
            "dataset_id": "out_2122-8th",
            "subject_code": "CSC801",
            "department": "Artificial Intelligence and Data Science",
            "batch": "2021–2022",
            "semester": "Semester VIII"
        }
        res = self.app.post("/api/subject-analysis/export-docx", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertIn("attachment;", res.headers.get("Content-Disposition", ""))
        self.assertIn(".docx", res.headers.get("Content-Disposition", ""))
        self.assertGreater(len(res.data), 1000)


if __name__ == "__main__":
    unittest.main()

