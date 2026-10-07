"""
Subject Performance Analysis Word (.docx) Exporter
===================================================
Generates a strictly 1-page institutional academic report in Microsoft Word format
for Terna Engineering College, complying with departmental audit guidelines.

Layout Architecture:
- Calibrated 1-page A4 format with institutional margins spanning the whole page.
- Top Left: Official Terna Engineering College Logo.
- Top Center: Institutional header hierarchy.
- Section 1: Theory Performance Summary Table + Clearance Donut Chart.
- Section 2: Laboratory Evaluation (Term Work & Practical / Oral) + Distribution Histograms.
- Section 3: Subject Toppers full-width structured merit table (no graph).
- Section 4: Institutional signatures block.
"""

import os
from io import BytesIO
from datetime import datetime
from typing import Dict, Any, Optional, List

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


# Institutional Color Palette
COLOR_NAVY = RGBColor(31, 78, 120)       # #1F4E78
COLOR_SLATE = RGBColor(51, 65, 85)       # #334155
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B
COLOR_SUCCESS = RGBColor(16, 185, 129)   # #10B981
COLOR_DANGER = RGBColor(239, 68, 68)     # #EF4444
HEX_NAVY = "1F4E78"
HEX_LIGHT_BLUE = "F0F4F8"
HEX_BORDER = "CBD5E1"
HEX_ZEBRA = "F8FAFC"


def _get_terna_logo_path() -> Optional[str]:
    """Finds the Terna logo file path in the project."""
    current_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        os.path.join(current_dir, "..", "static", "img", "terna_logo.png"),
        os.path.join(current_dir, "..", "static", "terna_logo.png"),
        os.path.join("d:\\Projects\\result-analysis-aids\\data_extractor_AI\\static\\img\\terna_logo.png")
    ]
    for p in candidates:
        norm = os.path.normpath(p)
        if os.path.exists(norm):
            return norm
    return None


def _set_cell_background(cell, hex_color: str):
    """Sets background shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top: int = 24, bottom: int = 24, left: int = 40, right: int = 40):
    """Sets internal padding (in dxa) for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_table_borders(table, color_hex: str = HEX_BORDER):
    """Applies neat institutional borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


# =========================================================================
# MATPLOTLIB CHART GENERATORS FOR EMBEDDING IN WORD DOCUMENT
# =========================================================================

def _generate_theory_chart_image(th: Dict[str, Any]) -> BytesIO:
    """Generates an uncluttered, compact clearance donut chart with structured legend for Theory summary."""
    passed = th.get("passed", 0)
    failed = th.get("failed", 0)
    first_class = th.get("first_class_count", 0)
    distinction = th.get("distinction_count", 0)
    other_pass = max(0, passed - (first_class + distinction))

    labels: List[str] = []
    sizes: List[float] = []
    colors: List[str] = []

    if distinction > 0:
        labels.append(f"Distinction ({distinction})")
        sizes.append(distinction)
        colors.append("#7c3aed")
    if first_class > 0:
        labels.append(f"1st Class ({first_class})")
        sizes.append(first_class)
        colors.append("#2563eb")
    if other_pass > 0:
        labels.append(f"Pass Class ({other_pass})")
        sizes.append(other_pass)
        colors.append("#10b981")
    if failed > 0:
        labels.append(f"Failed ({failed})")
        sizes.append(failed)
        colors.append("#ef4444")

    if not sizes:
        sizes = [1]
        labels = ["No Data"]
        colors = ["#cbd5e1"]

    fig, ax = plt.subplots(figsize=(2.60, 1.75), dpi=180)
    wedges, texts, autotexts = ax.pie(
        sizes,
        colors=colors,
        autopct="%1.0f%%" if (len(sizes) > 1 or sizes[0] > 1) else "",
        pctdistance=0.72,
        startangle=140,
        textprops={"fontsize": 6.2, "color": "white", "fontweight": "bold"}
    )

    # Donut hole
    centre_circle = plt.Circle((0, 0), 0.50, fc="white")
    ax.add_artist(centre_circle)
    ax.axis("equal")

    # Clean multi-column legend below the donut
    ax.legend(
        wedges, labels,
        loc="lower center",
        bbox_to_anchor=(0.5, -0.25),
        ncol=2,
        fontsize=6.0,
        frameon=False,
        handletextpad=0.3,
        columnspacing=0.8
    )

    ax.set_title("Clearance Breakdown", fontsize=7.5, fontweight="bold", color="#1F4E78", pad=2)
    plt.tight_layout()

    buf = BytesIO()
    plt.savefig(buf, format="png", bbox_inches="tight", transparent=False)
    plt.close(fig)
    buf.seek(0)
    return buf


def _generate_bar_chart_image(chart_info: Dict[str, Any], title: str, color_hex: str = "#2563eb") -> BytesIO:
    """Generates an uncluttered, compact marks distribution bar chart for TW or PR/OR."""
    labels = chart_info.get("labels", [])
    counts = chart_info.get("counts", [])

    fig, ax = plt.subplots(figsize=(3.15, 1.25), dpi=180)
    if labels and counts and any(c > 0 for c in counts):
        clean_labels = [str(l).replace("%", "") for l in labels]
        bars = ax.bar(clean_labels, counts, color=color_hex, edgecolor="#1e293b", linewidth=0.35, width=0.52)
        ax.set_ylabel("Students", fontsize=6.0, color="#475569", labelpad=2)
        ax.tick_params(axis="both", labelsize=6.2, pad=1)
        max_c = max(counts) if counts else 1
        ax.set_ylim(0, max_c * 1.25)
        for b, c in zip(bars, counts):
            if c > 0:
                ax.text(
                    b.get_x() + b.get_width() / 2,
                    b.get_height() + (max_c * 0.04),
                    str(c),
                    ha="center",
                    va="bottom",
                    fontsize=6.0,
                    fontweight="bold",
                    color="#1e293b"
                )
        ax.grid(axis="y", linestyle="--", alpha=0.3, color="#94a3b8")
    else:
        ax.text(0.5, 0.5, "Evaluation Pending / N/A", ha="center", va="center", fontsize=7.2, color="#64748b")
        ax.axis("off")

    ax.set_title(title, fontsize=7.5, fontweight="bold", color="#1F4E78", pad=2)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#cbd5e1")
    ax.spines["bottom"].set_color("#cbd5e1")
    plt.tight_layout()

    buf = BytesIO()
    plt.savefig(buf, format="png", bbox_inches="tight", transparent=False)
    plt.close(fig)
    buf.seek(0)
    return buf


def generate_subject_analysis_docx(
    analysis_data: Dict[str, Any],
    department_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    semester_name: Optional[str] = None
) -> BytesIO:
    """
    Creates a full-span, strictly 1-page institutional .docx document from subject analysis data.
    Neat, professional, well-utilized page space, and free of clutter.
    """
    doc = docx.Document()

    # Page Margins - Calibrated to span full A4 without spillover
    for section in doc.sections:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69) # A4

    # Extract Data
    hdr = analysis_data.get("header", {})
    th = analysis_data.get("theory", {})
    tw = analysis_data.get("term_work", {})
    pr = analysis_data.get("practical_oral", {}) or analysis_data.get("practical", {})
    toppers = analysis_data.get("top_students", [])

    dept_display = department_name or hdr.get("branch") or hdr.get("department") or "Engineering"
    if not dept_display.lower().startswith("department"):
        dept_display = f"Department of {dept_display.replace('_', ' ')}"
    else:
        dept_display = dept_display.replace("_", " ")

    batch_display = batch_name or hdr.get("batch") or hdr.get("academic_year") or "2021–2022"
    sem_display = semester_name or hdr.get("semester") or "Semester VIII"
    sem_display = str(sem_display).replace("Semester_", "Semester ").replace("Sem_", "Semester ")

    subj_code = hdr.get("subject_code") or "COURSE"
    subj_name = hdr.get("subject_name") or "Subject Analysis"
    max_marks = float(hdr.get("total_max_marks") or 100)

    # =========================================================================
    # HEADER SECTION WITH TOP-LEFT LOGO & CENTERED TITLES
    # =========================================================================
    logo_path = _get_terna_logo_path()
    tbl_hdr = doc.add_table(rows=1, cols=3)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hdr.autofit = False
    col_w_hdr = [Inches(1.10), Inches(5.10), Inches(1.10)]
    row_hdr = tbl_hdr.rows[0]

    # Col 0: Terna Logo (Top Left)
    c_logo = row_hdr.cells[0]
    c_logo.width = col_w_hdr[0]
    _set_cell_margins(c_logo, top=0, bottom=0, left=0, right=10)
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(0)
    if logo_path and os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, width=Inches(0.92))

    # Col 1: Institutional Header Text (Centered)
    c_text = row_hdr.cells[1]
    c_text.width = col_w_hdr[1]
    _set_cell_margins(c_text, top=0, bottom=0, left=5, right=5)

    p_inst = c_text.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(1)
    p_inst.paragraph_format.line_spacing = Pt(16)
    run_inst = p_inst.add_run("TERNA ENGINEERING COLLEGE")
    run_inst.bold = True
    run_inst.font.size = Pt(15.5)
    run_inst.font.color.rgb = COLOR_NAVY
    run_inst.font.name = "Arial"

    p_dept = c_text.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(0)
    p_dept.paragraph_format.space_after = Pt(1)
    p_dept.paragraph_format.line_spacing = Pt(12)
    run_dept = p_dept.add_run(dept_display.upper())
    run_dept.bold = True
    run_dept.font.size = Pt(10.5)
    run_dept.font.color.rgb = COLOR_SLATE
    run_dept.font.name = "Arial"

    p_title = c_text.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(0)
    p_title.paragraph_format.line_spacing = Pt(10)
    run_title = p_title.add_run("SUBJECT EXAMINATION PERFORMANCE & CLEARANCE REPORT")
    run_title.bold = True
    run_title.font.size = Pt(8.5)
    run_title.font.color.rgb = COLOR_MUTED
    run_title.font.name = "Arial"

    # Col 2: Balancing spacer cell
    c_spc = row_hdr.cells[2]
    c_spc.width = col_w_hdr[2]
    _set_cell_margins(c_spc, top=0, bottom=0, left=0, right=0)

    # Divider bar
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(4)
    p_div.paragraph_format.line_spacing = Pt(4)
    run_div = p_div.add_run("―" * 68)
    run_div.font.size = Pt(5)
    run_div.font.color.rgb = RGBColor(180, 195, 215)

    # =========================================================================
    # METADATA BLOCK
    # =========================================================================
    tbl_meta = doc.add_table(rows=3, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_meta.autofit = False

    meta_items = [
        ("Batch:", f"{batch_display}", "Academic Year:", f"{hdr.get('academic_year', batch_display)}"),
        ("Semester:", f"{sem_display}", "Evaluation Max:", f"{int(max_marks) if max_marks.is_integer() else max_marks} Marks"),
        ("Subject:", f"{subj_code} - {subj_name}", "Report Date:", datetime.now().strftime("%d-%b-%Y"))
    ]

    col_widths_meta = [Inches(5.0), Inches(2.30)]
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_items):
        row = tbl_meta.rows[r_idx]

        # Left cell
        cell_l = row.cells[0]
        cell_l.width = col_widths_meta[0]
        _set_cell_margins(cell_l, top=14, bottom=14, left=0, right=15)
        p_l = cell_l.paragraphs[0]
        p_l.paragraph_format.space_before = Pt(0)
        p_l.paragraph_format.space_after = Pt(0)
        p_l.paragraph_format.line_spacing = Pt(11)
        r_k1 = p_l.add_run(f"{k1} ")
        r_k1.bold = True
        r_k1.font.size = Pt(9.2)
        r_k1.font.color.rgb = COLOR_SLATE
        r_v1 = p_l.add_run(v1)
        r_v1.bold = (k1 == "Subject:")
        r_v1.font.size = Pt(9.2)
        if k1 == "Subject:":
            r_v1.font.color.rgb = COLOR_NAVY

        # Right cell
        cell_r = row.cells[1]
        cell_r.width = col_widths_meta[1]
        _set_cell_margins(cell_r, top=14, bottom=14, left=15, right=0)
        p_r = cell_r.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_r.paragraph_format.space_before = Pt(0)
        p_r.paragraph_format.space_after = Pt(0)
        p_r.paragraph_format.line_spacing = Pt(11)
        r_k2 = p_r.add_run(f"{k2} ")
        r_k2.bold = True
        r_k2.font.size = Pt(8.8)
        r_k2.font.color.rgb = COLOR_MUTED
        r_v2 = p_r.add_run(v2)
        r_v2.font.size = Pt(8.8)
        r_v2.font.color.rgb = COLOR_SLATE

    # =========================================================================
    # 1. THEORY EXAMINATION & CLEARANCE SUMMARY (TABLE + EMBEDDED GRAPH)
    # =========================================================================
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(4)
    p_sec1.paragraph_format.space_after = Pt(2)
    p_sec1.paragraph_format.line_spacing = Pt(11)
    r_s1 = p_sec1.add_run("1. THEORY EXAMINATION & CLEARANCE ANALYSIS")
    r_s1.bold = True
    r_s1.font.size = Pt(9.5)
    r_s1.font.color.rgb = COLOR_NAVY

    appeared = th.get("appeared", 0)
    passed = th.get("passed", 0)
    failed = th.get("failed", 0)
    pass_rate = th.get("pass_rate", 0.0)
    fail_rate = th.get("fail_rate", 0.0)

    first_class_cnt = th.get("first_class_count", 0)
    first_class_pct = th.get("first_class_rate", 0.0)
    dist_cnt = th.get("distinction_count", 0)
    dist_pct = th.get("distinction_rate", 0.0)

    fc_and_dist_cnt = th.get("first_class_and_dist_count", first_class_cnt + dist_cnt)
    fc_and_dist_pct = th.get("first_class_and_dist_rate", 0.0) or th.get("first_class_and_dist_pct", 0.0)

    table_data_th = [
        ("Total Students Appeared", appeared, "100.00%", None),
        ("Students Passed (≥ Passing Threshold)", passed, f"{pass_rate:.2f}%", COLOR_SUCCESS),
        ("Students Failed (Backlog / KT)", failed, f"{fail_rate:.2f}%", COLOR_DANGER),
        ("Students with First Class (60% to 74.99%)", first_class_cnt, f"{first_class_pct:.2f}%", None),
        ("Students with Distinction (≥ 75%)", dist_cnt, f"{dist_pct:.2f}%", None),
        ("Total First Class & Distinction (≥ 60%)", fc_and_dist_cnt, f"{fc_and_dist_pct:.2f}%", COLOR_NAVY),
    ]

    tbl_th = doc.add_table(rows=len(table_data_th) + 1, cols=4)
    tbl_th.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(tbl_th)
    col_w_th = [Inches(2.85), Inches(0.85), Inches(1.00), Inches(2.60)]

    # Header Row
    hdr_row_th = tbl_th.rows[0]
    th_headers = ["Clearance Status / Metric", "Count", "Percentage", "Clearance Graph"]
    for i, h_text in enumerate(th_headers):
        c = hdr_row_th.cells[i]
        c.width = col_w_th[i]
        _set_cell_background(c, HEX_NAVY)
        _set_cell_margins(c, top=32, bottom=32, left=45, right=45)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 3 else (WD_ALIGN_PARAGRAPH.RIGHT if i > 0 else WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.2)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, (label, count, pct, text_color) in enumerate(table_data_th):
        row = tbl_th.rows[r_idx + 1]
        is_highlight = ("Total First Class" in label or "Students Passed" in label)
        bg = HEX_LIGHT_BLUE if is_highlight else (HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF")

        for c_idx, val in enumerate([label, str(count), str(pct)]):
            cell = row.cells[c_idx]
            cell.width = col_w_th[c_idx]
            _set_cell_background(cell, bg)
            _set_cell_margins(cell, top=22, bottom=22, left=45, right=45)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.size = Pt(8.2)
            if is_highlight or c_idx > 0:
                run.bold = True
            if text_color:
                run.font.color.rgb = text_color
            else:
                run.font.color.rgb = COLOR_SLATE

    # Merge Col 3 across rows 1 to 6 for the Donut Chart
    cell_top_chart = tbl_th.cell(1, 3)
    cell_bot_chart = tbl_th.cell(len(table_data_th), 3)
    merged_th_chart = cell_top_chart.merge(cell_bot_chart)
    merged_th_chart.width = col_w_th[3]
    _set_cell_margins(merged_th_chart, top=8, bottom=8, left=15, right=15)
    _set_cell_background(merged_th_chart, "FFFFFF")
    th_chart_buf = _generate_theory_chart_image(th)
    p_c_th = merged_th_chart.paragraphs[0]
    p_c_th.paragraph_format.space_before = Pt(0)
    p_c_th.paragraph_format.space_after = Pt(0)
    p_c_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c_th.add_run().add_picture(th_chart_buf, width=Inches(2.55))

    # =========================================================================
    # 2. LABORATORY EVALUATION (TERM WORK & PRACTICAL WITH DUAL CHARTS)
    # =========================================================================
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(4)
    p_sec2.paragraph_format.space_after = Pt(2)
    p_sec2.paragraph_format.line_spacing = Pt(11)
    r_s2 = p_sec2.add_run("2. LABORATORY EVALUATION (TERM WORK & PRACTICAL / ORAL)")
    r_s2.bold = True
    r_s2.font.size = Pt(9.5)
    r_s2.font.color.rgb = COLOR_NAVY

    tw_avail = tw.get("available", False)
    pr_avail = pr.get("available", False)

    tbl_lab = doc.add_table(rows=6, cols=4)
    tbl_lab.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(tbl_lab)
    col_w_lab = [Inches(2.40), Inches(1.25), Inches(2.40), Inches(1.25)]

    hdr_lab = tbl_lab.rows[0]
    lab_headers = ["Term Work (TW) Performance", "Metric", "Practical / Oral (PR / OR)", "Metric"]
    for i, h_text in enumerate(lab_headers):
        c = hdr_lab.cells[i]
        c.width = col_w_lab[i]
        _set_cell_background(c, "2C5282")
        _set_cell_margins(c, top=32, bottom=32, left=40, right=40)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (i % 2 == 1) else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.2)
        r.font.color.rgb = RGBColor(255, 255, 255)

    tw_rows = [
        ("TW Students Appeared", f"{tw.get('appeared', 0)}" if tw_avail else "N/A"),
        ("TW Students Passed", f"{tw.get('passed', 0)}" if tw_avail else "N/A"),
        ("TW Students Failed", f"{tw.get('failed', 0)}" if tw_avail else "N/A"),
        ("TW Pass Percentage", f"{tw.get('pass_rate', 0.0):.2f}%" if tw_avail else "N/A"),
    ]

    pr_rows = [
        ("PR/OR Students Appeared", f"{pr.get('appeared', 0)}" if pr_avail else "N/A"),
        ("PR/OR Students Passed", f"{pr.get('passed', 0)}" if pr_avail else "N/A"),
        ("PR/OR Students Failed", f"{pr.get('failed', 0)}" if pr_avail else "N/A"),
        ("PR/OR Pass Percentage", f"{pr.get('pass_rate', 0.0):.2f}%" if pr_avail else "N/A"),
    ]

    for r_idx in range(4):
        row = tbl_lab.rows[r_idx + 1]
        bg = HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF"

        # TW
        tw_lbl, tw_val = tw_rows[r_idx]
        for c_i, text in [(0, tw_lbl), (1, tw_val)]:
            c = row.cells[c_i]
            c.width = col_w_lab[c_i]
            _set_cell_background(c, bg)
            _set_cell_margins(c, top=20, bottom=20, left=40, right=40)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_i == 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(8.2)
            if c_i == 1:
                run.bold = True

        # PR/OR
        pr_lbl, pr_val = pr_rows[r_idx]
        for c_i, text in [(2, pr_lbl), (3, pr_val)]:
            c = row.cells[c_i]
            c.width = col_w_lab[c_i]
            _set_cell_background(c, bg)
            _set_cell_margins(c, top=20, bottom=20, left=40, right=40)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_i == 3 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(8.2)
            if c_i == 3:
                run.bold = True

    # Row 5: Charts row
    row_charts = tbl_lab.rows[5]
    c_tw_chart = row_charts.cells[0].merge(row_charts.cells[1])
    c_pr_chart = row_charts.cells[2].merge(row_charts.cells[3])
    _set_cell_margins(c_tw_chart, top=10, bottom=10, left=15, right=15)
    _set_cell_margins(c_pr_chart, top=10, bottom=10, left=15, right=15)
    _set_cell_background(c_tw_chart, "FFFFFF")
    _set_cell_background(c_pr_chart, "FFFFFF")

    tw_chart_buf = _generate_bar_chart_image(tw.get("marks_distribution", {}).get("chart", {}), "TW Score Distribution", "#a855f7")
    p_tw_c = c_tw_chart.paragraphs[0]
    p_tw_c.paragraph_format.space_before = Pt(0)
    p_tw_c.paragraph_format.space_after = Pt(0)
    p_tw_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tw_c.add_run().add_picture(tw_chart_buf, width=Inches(3.10))

    pr_chart_buf = _generate_bar_chart_image(pr.get("marks_distribution", {}).get("chart", {}), "PR/OR Score Distribution", "#14b8a6")
    p_pr_c = c_pr_chart.paragraphs[0]
    p_pr_c.paragraph_format.space_before = Pt(0)
    p_pr_c.paragraph_format.space_after = Pt(0)
    p_pr_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pr_c.add_run().add_picture(pr_chart_buf, width=Inches(3.10))

    # =========================================================================
    # 3. SUBJECT TOPPERS (FULL-WIDTH STRUCTURED TABLE - NO GRAPH)
    # =========================================================================
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(4)
    p_sec3.paragraph_format.space_after = Pt(2)
    p_sec3.paragraph_format.line_spacing = Pt(11)
    r_s3 = p_sec3.add_run("3. TOP 5 MERIT LEADERS (COMBINED SUBJECT RESULT: THEORY + LAB)")
    r_s3.bold = True
    r_s3.font.size = Pt(9.5)
    r_s3.font.color.rgb = COLOR_NAVY

    toppers_to_show = toppers[:5]
    row_count_top = max(1, len(toppers_to_show)) + 1

    tbl_top = doc.add_table(rows=row_count_top, cols=5)
    tbl_top.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(tbl_top)
    col_w_top = [Inches(0.85), Inches(2.70), Inches(1.40), Inches(1.15), Inches(1.20)]

    # Header Row
    hdr_top = tbl_top.rows[0]
    top_headers = ["Rank", "Student Name", "Seat No / PRN", "Total Marks", "Percentage"]
    for i, h_text in enumerate(top_headers):
        c = hdr_top.cells[i]
        c.width = col_w_top[i]
        _set_cell_background(c, HEX_NAVY)
        _set_cell_margins(c, top=32, bottom=32, left=40, right=40)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else (WD_ALIGN_PARAGRAPH.RIGHT if i in (3, 4) else WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.2)
        r.font.color.rgb = RGBColor(255, 255, 255)

    if not toppers_to_show:
        empty_row = tbl_top.rows[1]
        for c in empty_row.cells:
            c.width = col_w_top[0]
            _set_cell_margins(c, top=24, bottom=24, left=40, right=40)
        p = empty_row.cells[1].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.add_run("No topper records available.").font.color.rgb = COLOR_MUTED
    else:
        medals = {1: "Rank 1 [🥇]", 2: "Rank 2 [🥈]", 3: "Rank 3 [🥉]", 4: "Rank 4", 5: "Rank 5"}
        for r_idx, st in enumerate(toppers_to_show):
            row = tbl_top.rows[r_idx + 1]
            bg = HEX_LIGHT_BLUE if r_idx == 0 else (HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF")
            rank_val = medals.get(st.get("rank"), f"Rank {st.get('rank', r_idx + 1)}")
            name_val = str(st.get("student_name", "")).strip().upper()
            seat_val = str(st.get("seat_or_prn", "")).strip()
            total_score_val = f"{st.get('marks', 0)} / {int(max_marks) if max_marks.is_integer() else max_marks}"
            pct_val = f"{st.get('percentage', 0.0):.2f}%"

            values = [rank_val, name_val, seat_val, total_score_val, pct_val]
            for c_idx, val in enumerate(values):
                cell = row.cells[c_idx]
                cell.width = col_w_top[c_idx]
                _set_cell_background(cell, bg)
                _set_cell_margins(cell, top=24, bottom=24, left=40, right=40)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(10)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx in (3, 4) else (WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT)
                run = p.add_run(val)
                run.font.size = Pt(8.2)
                if c_idx == 0 or r_idx == 0:
                    run.bold = True
                if c_idx == 3:
                    run.bold = True
                    run.font.color.rgb = COLOR_NAVY
                elif c_idx == 4:
                    run.bold = True
                    run.font.color.rgb = COLOR_SUCCESS
                else:
                    run.font.color.rgb = COLOR_SLATE

    # =========================================================================
    # 4. INSTITUTIONAL SIGNATURES BLOCK
    # =========================================================================
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(14)
    p_sp2.paragraph_format.space_after = Pt(0)
    p_sp2.paragraph_format.line_spacing = Pt(1)

    tbl_sig = doc.add_table(rows=1, cols=3)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_sig = [Inches(2.45), Inches(2.40), Inches(2.45)]

    sig_titles = [
        "Subject Faculty In-Charge",
        "Academic Coordinator",
        "Head of Department (HOD)"
    ]

    for i, title in enumerate(sig_titles):
        c = tbl_sig.rows[0].cells[i]
        c.width = col_widths_sig[i]
        _set_cell_margins(c, top=0, bottom=0, left=20, right=20)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(10)

        r_line = p.add_run("__________________________\n")
        r_line.font.size = Pt(8.2)
        r_line.font.color.rgb = COLOR_MUTED

        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.size = Pt(8.2)
        r_t.font.color.rgb = COLOR_SLATE

        r_dept = p.add_run(f"\n{dept_display}")
        r_dept.font.size = Pt(7.2)
        r_dept.font.color.rgb = COLOR_MUTED

    # Save to buffer
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
