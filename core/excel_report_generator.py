"""
Excel Academic Audit & Result Analysis Report Generator
Replicates the exact 6-sheet university departmental analysis workbook:
  1. Data
  2. IAT KT analysis
  3. PR KT analysis
  4. Summary
  5. Subject Wise Result
  6. Distributions

Outputs an Excel file formatted precisely as requested (e.g. "21-22 sem 8.xlsx").
"""

import os
import re
import json
import logging
from typing import List, Dict, Any, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# --- Color & Style Palette ---
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

SECTION_TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1F4E78")
SECTION_SUBTITLE_FONT = Font(name="Calibri", size=11, bold=True, color="333333")

TABLE_HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TABLE_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="000000")

HIGHLIGHT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
PASS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

DATA_FONT = Font(name="Calibri", size=11, bold=False, color="000000")
BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="000000")

THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)

HEADER_BORDER = Border(
    left=Side(style="thin", color="FFFFFF"),
    right=Side(style="thin", color="FFFFFF"),
    top=Side(style="medium", color="1F4E78"),
    bottom=Side(style="medium", color="1F4E78"),
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _safe_float(val: Any) -> Optional[float]:
    """Safely convert value to float or return None."""
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in ["", "nan", "null", "none", "aa", "ab", "abs", "--", "*"]:
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _is_absent(val: Any) -> bool:
    """Check if mark or grade indicates student was absent."""
    if val is None:
        return False
    s = str(val).strip().upper()
    return s in ["AA", "AB", "ABS", "ABSENT", "A"]


def _get_est_max_marks(max_val: Optional[float], default: float = 100.0) -> float:
    """Infer course total max marks ceiling from highest marks scored."""
    if not max_val or max_val <= 0:
        return default
    for threshold in [25.0, 50.0, 75.0, 100.0, 125.0, 150.0, 175.0, 200.0]:
        if max_val <= threshold:
            return threshold
    return default


def get_academic_filename(folder_name: str, meta: Optional[Dict[str, Any]] = None) -> str:
    """
    Produce standard academic file title matching target format (e.g. '21-22 sem 8.xlsx').
    """
    m = meta or {}
    year = str(m.get("academic_year") or "").strip()
    sem = str(m.get("semester") or "").strip()

    year_str = ""
    # Check folder name first, e.g. 2122-8th or 2223-7th
    fn_m = re.match(r"^(\d{2})(\d{2})[-_](\w+)", folder_name)
    if fn_m:
        year_str = f"{fn_m.group(1)}-{fn_m.group(2)}"
        if not sem:
            sem = fn_m.group(3)
    elif year:
        y_m = re.search(r"(\d{2})(\d{2})[-_](\d{2,4})", year) or re.search(r"(\d{2})[-_](\d{2})", year)
        if y_m:
            if len(y_m.groups()) >= 3 and y_m.group(3):
                y2 = y_m.group(3)[-2:]
                year_str = f"{y_m.group(2)}-{y2}"
            else:
                year_str = f"{y_m.group(1)}-{y_m.group(2)}"

    if not year_str:
        year_str = folder_name or "academic"

    # Normalize semester (VIII -> 8, VII -> 7, 8th -> 8)
    roman_map = {"I": "1", "II": "2", "III": "3", "IV": "4", "V": "5", "VI": "6", "VII": "7", "VIII": "8"}
    sem_clean = sem.upper().replace("SEM", "").replace("SEMESTER", "").strip()
    if sem_clean in roman_map:
        sem_str = roman_map[sem_clean]
    else:
        num_m = re.search(r"(\d+)", sem or folder_name)
        sem_str = num_m.group(1) if num_m else "1"

    return f"{year_str} sem {sem_str}.xlsx"


HEAD_ALIASES: Dict[str, List[str]] = {
    "TH": ["TH", "SE", "THEORY", "EXT", "EXTERNAL", "ESE", "SEM_EXAM"],
    "IA": ["IA", "IN", "IAT", "INTERNAL", "INT", "TEST", "CA", "ISA"],
    "PR": ["PR", "OR", "PR_OR", "OR_PR", "TW_PR", "PRACTICAL", "ORAL", "VIVA", "LAB", "PO"],
    "TW": ["TW", "TERM_WORK", "TERMWORK", "TERM WORK"],
    "TOT": ["TOT", "TOTAL", "TOT_MARKS", "MARKS"]
}

NON_COURSE_PREFIXES = {
    "SEAT", "STUDENT", "CANDIDATE", "TOTAL", "CREDITS", "RESULT",
    "MOTHER", "FATHER", "ABC", "PRN", "ERN", "REMARK", "SGPI", "CGPI"
}


class AcademicExcelReportGenerator:
    """
    Constructs a publication-grade, NAAC/NBA compliant 6-sheet analysis Excel workbook.
    """

    def __init__(
        self,
        students: List[Dict[str, Any]],
        catalog: Optional[Any] = None,
        meta: Optional[Dict[str, Any]] = None,
        folder_name: str = "",
        fieldnames: Optional[List[str]] = None
    ):
        self.catalog = catalog or {}
        self.meta = meta or {}
        self.folder_name = folder_name or "report"
        self.fieldnames = list(fieldnames) if fieldnames else None

        # Normalize students to list of flat dictionaries
        self.students = []
        for s in (students or []):
            if hasattr(s, "to_flat_dict"):
                self.students.append(s.to_flat_dict(self.catalog))
            elif isinstance(s, dict):
                self.students.append(s)
            else:
                self.students.append(dict(s))

        self.courses = self._resolve_courses()

    def _resolve_courses(self) -> List[Dict[str, Any]]:
        """Identify all courses and their heads from catalog or dataframe columns."""
        cat_courses = []
        if hasattr(self.catalog, "to_dict"):
            try:
                cat_courses = self.catalog.to_dict()
            except Exception:
                cat_courses = []
        elif isinstance(self.catalog, dict):
            cat_courses = self.catalog.get("courses", [])
        elif isinstance(self.catalog, list):
            cat_courses = self.catalog

        resolved = []
        if cat_courses:
            for c in cat_courses:
                if hasattr(c, "code"):
                    resolved.append({
                        "code": str(getattr(c, "code", "")).strip().upper(),
                        "name": str(getattr(c, "name", "")).strip() or str(getattr(c, "code", "")).strip().upper(),
                        "heads": list(getattr(c, "heads", [])),
                        "max_marks": dict(getattr(c, "max_marks", {}) or {}),
                        "min_marks": dict(getattr(c, "min_marks", {}) or {}),
                        "credits": float(getattr(c, "credits", 0.0) or 0.0)
                    })
                elif isinstance(c, dict):
                    c_code = str(c.get("code", "")).strip().upper()
                    if c_code:
                        resolved.append({
                            "code": c_code,
                            "name": str(c.get("name", "")).strip() or c_code,
                            "heads": list(c.get("heads", [])),
                            "max_marks": dict(c.get("max_marks", {}) or {}),
                            "min_marks": dict(c.get("min_marks", {}) or {}),
                            "credits": float(c.get("credits", 0.0) or 0.0)
                        })
            if resolved:
                return resolved

        # Fallback: discover courses from student dict keys
        if not self.students:
            return []

        first = self.students[0]
        discovered = {}
        for k in first.keys():
            if "_" in k:
                parts = k.split("_", 1)
                prefix = parts[0]
                suffix = parts[1].upper()
                if prefix.upper() in NON_COURSE_PREFIXES:
                    continue
                valid_suffixes = [
                    "TH", "SE", "IN", "IA", "IAT", "TW", "PR", "OR", "OR_PR",
                    "PR_OR", "TW_PR", "TOTAL", "TOT", "GRADE", "GP", "CREDITS", "CG", "SUBJECT"
                ]
                if suffix in valid_suffixes:
                    if prefix not in discovered:
                        discovered[prefix] = {
                            "code": prefix,
                            "name": first.get(f"{prefix}_Subject", first.get(f"{prefix}_Name", prefix)),
                            "heads": [],
                            "max_marks": {},
                            "min_marks": {},
                            "credits": 0.0
                        }
                    head_suffixes = ["TH", "SE", "IN", "IA", "IAT", "TW", "PR", "OR", "OR_PR", "PR_OR", "TW_PR", "TOTAL", "TOT"]
                    if suffix in head_suffixes and suffix not in discovered[prefix]["heads"]:
                        discovered[prefix]["heads"].append(suffix)

        return list(discovered.values())

    def _get_student_mark(self, s: Dict[str, Any], course_code: str, head_type: str) -> Optional[float]:
        """Retrieve student's mark for a specific course head category using aliases."""
        aliases = HEAD_ALIASES.get(head_type, [head_type])
        for a in aliases:
            for pat in [f"{course_code}_{a}", f"{course_code}_{a.lower()}", f"{course_code}_{a.capitalize()}", f"{course_code}_{a.upper()}"]:
                if pat in s:
                    val = _safe_float(s[pat])
                    if val is not None:
                        return val
        return None

    def _is_student_absent_in_head(self, s: Dict[str, Any], course_code: str, head_type: str) -> bool:
        """Check if student was absent in a specific course head category."""
        aliases = HEAD_ALIASES.get(head_type, [head_type])
        for a in aliases:
            for pat in [f"{course_code}_{a}", f"{course_code}_{a.lower()}", f"{course_code}_{a.capitalize()}", f"{course_code}_{a.upper()}"]:
                if pat in s and _is_absent(s[pat]):
                    return True
        return False

    def _is_student_enrolled_in_course(self, s: Dict[str, Any], course_code: str) -> bool:
        """Return True if the student was enrolled in this course."""
        for suffix in ["_Total", "_TOT", "_TH", "_SE", "_IN", "_IA", "_TW", "_PR", "_OR", "_OR_PR", "_Grade"]:
            val = s.get(f"{course_code}{suffix}")
            if val is not None and str(val).strip() not in ["", "--", "nan", "NaN", "None"]:
                return True
        return False

    def _get_course_head_max(self, course: Dict[str, Any], head_type: str, default: Optional[float] = None) -> float:
        """Determine maximum marks for course head category."""
        raw_max_dict = course.get("max_marks", {}) or {}
        aliases = HEAD_ALIASES.get(head_type, [head_type])
        for a in aliases:
            for k, v in raw_max_dict.items():
                if k.upper() == a.upper():
                    val = _safe_float(v)
                    if val and val > 0:
                        return val

        # Estimate from enrolled students' highest mark
        max_scored = 0.0
        code = course.get("code", "")
        for s in self.students:
            m = self._get_student_mark(s, code, head_type)
            if m is not None and m > max_scored:
                max_scored = m

        if max_scored > 0:
            if head_type == "IA":
                return 40.0 if max_scored > 20.0 else 20.0
            elif head_type == "TH":
                return 80.0 if max_scored > 60.0 else 60.0
            elif head_type in ["PR", "TW"]:
                return 50.0 if max_scored > 25.0 else 25.0
            elif head_type == "TOT":
                return _get_est_max_marks(max_scored, default=100.0)

        if default is not None:
            return default

        # Scheme Defaults
        if head_type == "IA":
            return 20.0
        elif head_type == "TH":
            return 80.0
        elif head_type in ["PR", "TW"]:
            return 25.0
        elif head_type == "TOT":
            return 100.0
        return 100.0

    def _get_course_head_min(self, course: Dict[str, Any], head_type: str) -> float:
        """Determine minimum passing mark for course head category using scheme or 40% rule."""
        raw_min_dict = course.get("min_marks", {}) or {}
        aliases = HEAD_ALIASES.get(head_type, [head_type])
        for a in aliases:
            for k, v in raw_min_dict.items():
                if k.upper() == a.upper():
                    val = _safe_float(v)
                    if val and val > 0:
                        return val

        # 40% Statutory Passing Rule
        import math
        max_m = self._get_course_head_max(course, head_type)
        return float(math.ceil(0.40 * max_m))

    def _is_lab_or_practical(self, course: Dict[str, Any]) -> bool:
        """Determine if course is a Practical, Lab, Oral, or Project."""
        name = course.get("name", "").lower()
        code = course.get("code", "").upper()
        heads = [h.upper() for h in course.get("heads", [])]
        if "lab" in name or "laboratory" in name or "practical" in name or "lab" in code.lower() or code.endswith("L"):
            return True
        if any(h in heads for h in ["PR", "OR", "OR_PR", "PR_OR", "TW_PR"]) and "TH" not in heads and "SE" not in heads:
            return True
        return False

    def _get_theory_courses(self) -> List[Dict[str, Any]]:
        """Return courses that represent Theory subjects (strictly excluding labs/practicals)."""
        return [c for c in self.courses if not self._is_lab_or_practical(c)]

    def _get_lab_courses(self) -> List[Dict[str, Any]]:
        """Return courses that represent Practicals, Labs, Orals, or Projects."""
        return [c for c in self.courses if self._is_lab_or_practical(c)]

    def _get_enrolled_students(self, course_code: str) -> List[Dict[str, Any]]:
        """Return only students enrolled in this specific course."""
        return [s for s in self.students if self._is_student_enrolled_in_course(s, course_code)]

    def _is_student_absent_in_course(self, s: Dict[str, Any], course: Any) -> bool:
        """Check if student was absent in this course."""
        code = course["code"] if isinstance(course, dict) else str(course)
        gr = str(s.get(f"{code}_Grade", "")).strip().upper()
        if gr in ["AB", "AA", "ABS", "A", "ABSENT"]:
            return True

        has_any_data = False
        all_absent = True
        for head in ["TOT", "TH", "IA", "TW", "PR"]:
            val = self._get_student_mark(s, code, head)
            if val is not None:
                has_any_data = True
                all_absent = False
                break
            if self._is_student_absent_in_head(s, code, head):
                has_any_data = True

        if has_any_data and all_absent:
            return True
        return False

    def _is_student_failed_in_course(self, s: Dict[str, Any], course: Any) -> bool:
        """Check if student failed this course according to scheme rules and Grade."""
        code = course["code"] if isinstance(course, dict) else str(course)
        c_dict = course if isinstance(course, dict) else {"code": code}

        gr = str(s.get(f"{code}_Grade", "")).strip().upper()
        if gr in ["F", "FAIL"]:
            return True

        # Check total mark against min_marks
        tot = self._get_student_mark(s, code, "TOT")
        min_tot = self._get_course_head_min(c_dict, "TOT")
        if tot is not None and tot < min_tot:
            return True

        # Check individual component heads
        th = self._get_student_mark(s, code, "TH")
        if th is not None:
            min_th = self._get_course_head_min(c_dict, "TH")
            if th < min_th:
                return True

        ia = self._get_student_mark(s, code, "IA")
        if ia is not None:
            min_ia = self._get_course_head_min(c_dict, "IA")
            if ia < min_ia:
                return True

        pr = self._get_student_mark(s, code, "PR")
        if pr is not None:
            min_pr = self._get_course_head_min(c_dict, "PR")
            if pr < min_pr:
                return True

        tw = self._get_student_mark(s, code, "TW")
        if tw is not None:
            min_tw = self._get_course_head_min(c_dict, "TW")
            if tw < min_tw:
                return True

        return False

    def generate_workbook(self, output_path: Optional[str] = None) -> openpyxl.Workbook:
        """Generate all 6 sheets and optionally write to disk."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        self._build_data_sheet(wb)
        self._build_iat_kt_sheet(wb)
        self._build_pr_kt_sheet(wb)
        self._build_summary_sheet(wb)
        self._build_subject_wise_sheet(wb)
        self._build_distributions_sheet(wb)

        if output_path:
            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
            wb.save(output_path)
            logger.info("Academic analysis workbook saved: %s", output_path)

        return wb

    # =========================================================================
    # 1. SHEET: Data
    # =========================================================================
    def _build_data_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title="Data")
        ws.views.sheetView[0].showGridLines = True

        # Resolve raw fieldnames directly from extracted data
        columns = list(self.fieldnames) if self.fieldnames else []
        if not columns and self.students:
            seen = set()
            for s in self.students:
                for k in s.keys():
                    if k not in seen:
                        seen.add(k)
                        columns.append(k)

        if not columns:
            return

        # Write Header Row
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = HEADER_BORDER

        ws.row_dimensions[1].height = 28

        # Identifier/text attributes preserved strictly as strings to prevent truncation
        string_cols = {
            "seat_no", "seat no", "seat", "prn", "abc_id", "ern", "student_name",
            "name", "gender", "candidate_type", "result_status", "remark", "semester"
        }

        # Write Raw Data Rows directly from extracted dataset
        for r_idx, s in enumerate(self.students, 2):
            for c_idx, col_name in enumerate(columns, 1):
                raw_val = s.get(col_name)
                col_lower = str(col_name).lower()

                cell_val = raw_val
                if raw_val is None or str(raw_val).strip() in ["", "nan", "NaN", "None"]:
                    cell_val = None
                elif col_lower in string_cols or col_lower.endswith(("_subject", "_name", "_grade", "_status")):
                    cell_val = str(raw_val).strip()
                else:
                    f_val = _safe_float(raw_val)
                    if f_val is not None:
                        if f_val.is_integer() and "." not in str(raw_val):
                            cell_val = int(f_val)
                        else:
                            cell_val = f_val
                    elif _is_absent(raw_val):
                        cell_val = "AA"
                    else:
                        cell_val = str(raw_val).strip()

                cell = ws.cell(row=r_idx, column=c_idx, value=cell_val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER

                # Alignment
                if isinstance(cell_val, (int, float)):
                    cell.alignment = ALIGN_RIGHT
                elif col_lower in ["student_name", "name"] or col_lower.endswith("_subject"):
                    cell.alignment = ALIGN_LEFT
                else:
                    cell.alignment = ALIGN_CENTER

            ws.row_dimensions[r_idx].height = 20

        ws.freeze_panes = "A2"
        self._auto_fit_columns(ws)

    # =========================================================================
    # 2. SHEET: IAT KT analysis
    # =========================================================================
    def _build_iat_kt_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title="IAT KT analysis")
        ws.views.sheetView[0].showGridLines = True

        # Find theory courses that possess an internal assessment component
        th_courses = []
        for c in self._get_theory_courses():
            heads = [h.upper() for h in c.get("heads", [])]
            if any(h in heads for h in ["IA", "IN", "IAT", "INTERNAL"]):
                th_courses.append(c)
            elif any(self._get_student_mark(s, c["code"], "IA") is not None for s in self.students[:10]):
                th_courses.append(c)

        if not th_courses:
            th_courses = self._get_theory_courses()[:6]

        headers = ["Seat No", "Name"] + [c.get("name", c["code"]) for c in th_courses] + ["Total IAT KT"]

        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c_idx, value=h)
            cell.fill = TABLE_HEADER_FILL
            cell.font = TABLE_HEADER_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        ws.row_dimensions[1].height = 26

        curr_row = 2
        for s in self.students:
            seat = s.get("Seat_No") or s.get("seat_no")
            name = s.get("Student_Name") or s.get("student_name") or ""
            row_cells = []
            kt_count = 0

            for c in th_courses:
                c_code = c["code"]
                if not self._is_student_enrolled_in_course(s, c_code):
                    row_cells.append("\u2014")
                    continue

                f_in = self._get_student_mark(s, c_code, "IA")
                is_abs = self._is_student_absent_in_head(s, c_code, "IA")
                min_ia = self._get_course_head_min(c, "IA")

                if is_abs:
                    row_cells.append("KT (Absent)")
                    kt_count += 1
                elif f_in is not None and f_in < min_ia:
                    row_cells.append("KT (Fail)")
                    kt_count += 1
                else:
                    row_cells.append("\u2014")

            # Only list students with >= 1 IAT KT
            if kt_count > 0:
                vals = [seat, str(name).strip()] + row_cells + [kt_count]
                for col_idx, val in enumerate(vals, 1):
                    cell = ws.cell(row=curr_row, column=col_idx, value=val)
                    cell.font = DATA_FONT
                    cell.border = THIN_BORDER
                    if val in ["KT (Fail)", "KT (Absent)"]:
                        cell.fill = FAIL_FILL
                        cell.font = BOLD_FONT
                    if col_idx == 2:
                        cell.alignment = ALIGN_LEFT
                    else:
                        cell.alignment = ALIGN_CENTER
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        self._auto_fit_columns(ws)

    # =========================================================================
    # 3. SHEET: PR KT analysis
    # =========================================================================
    def _build_pr_kt_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title="PR KT analysis")
        ws.views.sheetView[0].showGridLines = True

        lab_courses = self._get_lab_courses()
        if not lab_courses:
            for c in self.courses:
                heads = [h.upper() for h in c.get("heads", [])]
                if any(h in heads for h in ["PR", "OR", "PR_OR", "OR_PR", "TW_PR"]):
                    if c not in lab_courses:
                        lab_courses.append(c)

        if not lab_courses:
            lab_courses = self.courses[-3:]  # fallback

        headers = ["Seat No", "Name"] + [f"{c.get('name', c['code'])} PR" for c in lab_courses] + ["Total PR KT"]

        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c_idx, value=h)
            cell.fill = TABLE_HEADER_FILL
            cell.font = TABLE_HEADER_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        ws.row_dimensions[1].height = 26

        curr_row = 2
        for s in self.students:
            seat = s.get("Seat_No") or s.get("seat_no")
            name = s.get("Student_Name") or s.get("student_name") or ""
            row_cells = []
            kt_count = 0

            for c in lab_courses:
                c_code = c["code"]
                if not self._is_student_enrolled_in_course(s, c_code):
                    row_cells.append("\u2014")
                    continue

                f_pr = self._get_student_mark(s, c_code, "PR")
                is_abs = self._is_student_absent_in_head(s, c_code, "PR") or str(s.get(f"{c_code}_Grade") or "").strip().upper() in ["AB", "AA", "ABS"]
                grade = str(s.get(f"{c_code}_Grade") or "").strip().upper()
                min_pr = self._get_course_head_min(c, "PR")

                if is_abs:
                    row_cells.append("KT (Absent)")
                    kt_count += 1
                elif grade in ["F", "FAIL"] or (f_pr is not None and f_pr < min_pr):
                    row_cells.append("KT (Fail)")
                    kt_count += 1
                else:
                    row_cells.append("\u2014")

            # Only list students with >= 1 PR KT
            if kt_count > 0:
                vals = [seat, str(name).strip()] + row_cells + [kt_count]
                for col_idx, val in enumerate(vals, 1):
                    cell = ws.cell(row=curr_row, column=col_idx, value=val)
                    cell.font = DATA_FONT
                    cell.border = THIN_BORDER
                    if val in ["KT (Fail)", "KT (Absent)"]:
                        cell.fill = FAIL_FILL
                        cell.font = BOLD_FONT
                    if col_idx == 2:
                        cell.alignment = ALIGN_LEFT
                    else:
                        cell.alignment = ALIGN_CENTER
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        self._auto_fit_columns(ws)

    # =========================================================================
    # 4. SHEET: Summary
    # =========================================================================
    def _build_summary_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title="Summary")
        ws.views.sheetView[0].showGridLines = True

        total_students = len(self.students)

        # 1. Section: Topper Students
        title_cell = ws.cell(row=1, column=1, value="Topper Students")
        title_cell.font = SECTION_TITLE_FONT

        t_headers = ["Sr. No.", "Seat No", "Name of the student", "Marks", "GPA"]
        for c_idx, h in enumerate(t_headers, 1):
            c = ws.cell(row=3, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        # Strictly filter passing candidates (toppers cannot have KT or failed subjects)
        passing_candidates = [
            s for s in self.students
            if str(s.get("Result_Status", "")).upper() not in ["F", "FAIL", "FAILED", "UNSUCCESSFUL", "ATKT", "ABS", "A", "ABSENT"]
            and (_safe_float(s.get("SGPI")) or 0.0) > 0.0
            and not any(self._is_student_failed_in_course(s, c) for c in self.courses if self._is_student_enrolled_in_course(s, c["code"]))
        ]
        if not passing_candidates:
            passing_candidates = self.students

        sorted_st = sorted(
            passing_candidates,
            key=lambda x: (_safe_float(x.get("SGPI")) or 0.0, _safe_float(x.get("Total_Marks")) or 0.0),
            reverse=True
        )

        for rank, s in enumerate(sorted_st[:5], 1):
            r_idx = 3 + rank
            seat = s.get("Seat_No") or s.get("seat_no")
            name = s.get("Student_Name") or s.get("student_name")
            marks = _safe_float(s.get("Total_Marks")) or 0.0
            gpa = _safe_float(s.get("SGPI")) or 0.0

            vals = [rank, seat, str(name).strip(), marks, gpa]
            for col_idx, val in enumerate(vals, 1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if col_idx == 3:
                    cell.alignment = ALIGN_LEFT
                elif col_idx in [1, 2]:
                    cell.alignment = ALIGN_CENTER
                else:
                    cell.alignment = ALIGN_RIGHT

        # 2. Section: Overall Result
        ws.cell(row=9, column=1, value="Overall Result").font = SECTION_TITLE_FONT

        o_headers = ["No. of Students", "ABS", "Pass (All Clear)", "Fail with KT", "Pass Percentage"]
        for c_idx, h in enumerate(o_headers, 1):
            c = ws.cell(row=11, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        abs_count = 0
        fail_count = 0
        pass_count = 0

        for s in self.students:
            res_stat = str(s.get("Result_Status", "")).strip().upper()
            sgpi = _safe_float(s.get("SGPI")) or 0.0

            # Absent check
            is_abs = res_stat in ["ABS", "A", "ABSENT"] or all(
                self._is_student_absent_in_course(s, c)
                for c in self.courses
                if self._is_student_enrolled_in_course(s, c["code"])
            )
            if is_abs and total_students > 0 and len(self.courses) > 0:
                abs_count += 1
                continue

            # Fail check
            is_fail = (
                res_stat in ["F", "FAIL", "FAILED", "UNSUCCESSFUL", "ATKT"]
                or sgpi == 0.0
                or any(self._is_student_failed_in_course(s, c) for c in self.courses if self._is_student_enrolled_in_course(s, c["code"]))
            )
            if is_fail:
                fail_count += 1
            else:
                pass_count += 1

        pass_pct = round(pass_count / total_students, 4) if total_students > 0 else 0.0

        o_vals = [total_students, abs_count, pass_count, fail_count, pass_pct]
        for col_idx, val in enumerate(o_vals, 1):
            cell = ws.cell(row=12, column=col_idx, value=val)
            cell.font = BOLD_FONT
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER
            if col_idx == 5:
                cell.number_format = "0.00%"

        # 3. Section: Grade Point Wise Result
        ws.cell(row=14, column=1, value="Grade Point Wise Result").font = SECTION_TITLE_FONT

        c_gt9 = sum(1 for s in self.students if (_safe_float(s.get("SGPI")) or 0.0) >= 9.0)
        c_8to9 = sum(1 for s in self.students if 8.0 <= (_safe_float(s.get("SGPI")) or 0.0) < 9.0)
        c_7to8 = sum(1 for s in self.students if 7.0 <= (_safe_float(s.get("SGPI")) or 0.0) < 8.0)
        c_fail = fail_count
        c_abs = abs_count

        g_headers = [None, ">9", "8-8.9", "7-7.9", "Fail", "Absent"]
        for c_idx, h in enumerate(g_headers, 1):
            if h:
                c = ws.cell(row=17, column=c_idx, value=h)
                c.fill = TABLE_HEADER_FILL
                c.font = TABLE_HEADER_FONT
                c.border = THIN_BORDER
                c.alignment = ALIGN_CENTER

        # Percentage Row
        pct_vals = [
            "Percentage",
            round(c_gt9 / total_students, 4) if total_students > 0 else 0.0,
            round(c_8to9 / total_students, 4) if total_students > 0 else 0.0,
            round(c_7to8 / total_students, 4) if total_students > 0 else 0.0,
            round(c_fail / total_students, 4) if total_students > 0 else 0.0,
            round(c_abs / total_students, 4) if total_students > 0 else 0.0
        ]
        for col_idx, val in enumerate(pct_vals, 1):
            cell = ws.cell(row=18, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER
            if col_idx > 1:
                cell.number_format = "0.00%"

        # Count Row
        cnt_vals = [total_students, c_gt9, c_8to9, c_7to8, c_fail, c_abs]
        for col_idx, val in enumerate(cnt_vals, 1):
            cell = ws.cell(row=19, column=col_idx, value=val)
            cell.font = BOLD_FONT
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER

        self._auto_fit_columns(ws)

    # =========================================================================
    # 5. SHEET: Subject Wise Result
    # =========================================================================
    def _build_subject_wise_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title="Subject Wise Result")
        ws.views.sheetView[0].showGridLines = True

        th_courses = self._get_theory_courses()
        lab_courses = self._get_lab_courses()

        # 1. Section: Subject wise Result
        ws.cell(row=3, column=1, value="Subject wise Result").font = SECTION_TITLE_FONT

        sub_headers = ["Subject", "Appeared", "Pass", "Absent", "Fail", "Pass Percentage"]
        for c_idx, h in enumerate(sub_headers, 1):
            c = ws.cell(row=5, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        row_idx = 6
        for c in th_courses:
            code = c["code"]
            name = c.get("name", code)

            enrolled = self._get_enrolled_students(code)
            if not enrolled:
                continue

            enrolled_cnt = len(enrolled)
            absent_cnt = sum(1 for s in enrolled if self._is_student_absent_in_course(s, c))
            appeared_cnt = max(0, enrolled_cnt - absent_cnt)
            fail_cnt = sum(1 for s in enrolled if not self._is_student_absent_in_course(s, c) and self._is_student_failed_in_course(s, c))
            pass_cnt = max(0, appeared_cnt - fail_cnt)

            pass_pct = round(pass_cnt / enrolled_cnt, 4) if enrolled_cnt > 0 else 0.0

            vals = [name, appeared_cnt, pass_cnt, absent_cnt, fail_cnt, pass_pct]
            for col_idx, val in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if col_idx == 1:
                    cell.alignment = ALIGN_LEFT
                elif col_idx == 6:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "0.00%"
                else:
                    cell.alignment = ALIGN_CENTER
            row_idx += 1

        # 2. Section: Subject Result Analysis (Percentage Ranges)
        row_idx += 2
        ws.cell(row=row_idx, column=1, value="Subject Result Analysis").font = SECTION_TITLE_FONT
        row_idx += 2

        analysis_headers = ["Subject", "Appeared", "90-100", "80-89", "70-79", "60-69", "50-59", "40-49", "Fail", "Absent"]
        for c_idx, h in enumerate(analysis_headers, 1):
            c = ws.cell(row=row_idx, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        row_idx += 1
        for c in th_courses:
            code = c["code"]
            name = c.get("name", code)

            enrolled = self._get_enrolled_students(code)
            if not enrolled:
                continue

            max_m = self._get_course_head_max(c, "TOT")

            buckets = {"90": 0, "80": 0, "70": 0, "60": 0, "50": 0, "40": 0, "fail": 0, "abs": 0}
            appeared = 0

            for s in enrolled:
                tot_val = self._get_student_mark(s, code, "TOT")
                if tot_val is None:
                    tot_val = self._get_student_mark(s, code, "TH")

                is_abs = self._is_student_absent_in_course(s, c)
                is_fail = self._is_student_failed_in_course(s, c)

                if is_abs:
                    buckets["abs"] += 1
                elif is_fail:
                    buckets["fail"] += 1
                    appeared += 1
                elif tot_val is not None and max_m > 0:
                    appeared += 1
                    pct = (tot_val / max_m) * 100.0
                    if pct >= 90:
                        buckets["90"] += 1
                    elif pct >= 80:
                        buckets["80"] += 1
                    elif pct >= 70:
                        buckets["70"] += 1
                    elif pct >= 60:
                        buckets["60"] += 1
                    elif pct >= 50:
                        buckets["50"] += 1
                    elif pct >= 40:
                        buckets["40"] += 1
                    else:
                        buckets["fail"] += 1
                else:
                    buckets["fail"] += 1
                    appeared += 1

            vals = [
                name, appeared, buckets["90"], buckets["80"], buckets["70"],
                buckets["60"], buckets["50"], buckets["40"], buckets["fail"], buckets["abs"]
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = ALIGN_LEFT if col_idx == 1 else ALIGN_CENTER
            row_idx += 1

        # 3. Section: Practical-wise Result
        row_idx += 2
        ws.cell(row=row_idx, column=1, value="Practical-wise Result").font = SECTION_TITLE_FONT
        row_idx += 2

        pr_headers = ["Lab", "Appeared", "Pass", "Fail", "Absent", "Pass Percentage"]
        for c_idx, h in enumerate(pr_headers, 1):
            c = ws.cell(row=row_idx, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        row_idx += 1
        for c in lab_courses:
            code = c["code"]
            name = c.get("name", code)
            if not any(w in name.lower() for w in ["lab", "practical", "project"]):
                display_name = f"{name} Practical"
            else:
                display_name = name

            enrolled = self._get_enrolled_students(code)
            if not enrolled:
                continue

            enrolled_cnt = len(enrolled)
            absent_cnt = sum(1 for s in enrolled if self._is_student_absent_in_course(s, c))
            appeared_cnt = max(0, enrolled_cnt - absent_cnt)
            fail_cnt = sum(1 for s in enrolled if not self._is_student_absent_in_course(s, c) and self._is_student_failed_in_course(s, c))
            pass_cnt = max(0, appeared_cnt - fail_cnt)

            pass_pct = round(pass_cnt / enrolled_cnt, 4) if enrolled_cnt > 0 else 0.0

            vals = [display_name, appeared_cnt, pass_cnt, fail_cnt, absent_cnt, pass_pct]
            for col_idx, val in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if col_idx == 1:
                    cell.alignment = ALIGN_LEFT
                elif col_idx == 6:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "0.00%"
                else:
                    cell.alignment = ALIGN_CENTER
            row_idx += 1

        # 4. Section: Subject wise Topper
        row_idx += 2
        ws.cell(row=row_idx, column=1, value="Subject wise Topper").font = SECTION_TITLE_FONT
        row_idx += 2

        top_headers = ["Subject", "Seat No.", "Name of the Student", "Marks", "%"]
        for c_idx, h in enumerate(top_headers, 1):
            c = ws.cell(row=row_idx, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        row_idx += 1
        for c in (th_courses + lab_courses):
            code = c["code"]
            name = c.get("name", code)

            enrolled = self._get_enrolled_students(code)
            if not enrolled:
                continue

            # Find passed student with highest marks in this subject
            top_scorer = None
            max_m = -1.0

            for s in enrolled:
                if self._is_student_failed_in_course(s, c):
                    continue
                tot_val = self._get_student_mark(s, code, "TOT")
                if tot_val is None:
                    tot_val = self._get_student_mark(s, code, "TH") or self._get_student_mark(s, code, "PR") or self._get_student_mark(s, code, "TW")
                if tot_val is not None and tot_val > max_m:
                    max_m = tot_val
                    top_scorer = s

            # Fallback if no student passed
            if not top_scorer:
                for s in enrolled:
                    tot_val = self._get_student_mark(s, code, "TOT") or self._get_student_mark(s, code, "TH")
                    if tot_val is not None and tot_val > max_m:
                        max_m = tot_val
                        top_scorer = s

            if top_scorer and max_m > 0:
                seat = top_scorer.get("Seat_No") or top_scorer.get("seat_no")
                s_name = top_scorer.get("Student_Name") or top_scorer.get("student_name")
                est_max = self._get_course_head_max(c, "TOT")

                pct = round(min((max_m / est_max) * 100.0, 100.0), 2) if est_max > 0 else 0.0

                vals = [name, seat, str(s_name).strip(), max_m, pct]
                for col_idx, val in enumerate(vals, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = DATA_FONT
                    cell.border = THIN_BORDER
                    if col_idx in [1, 3]:
                        cell.alignment = ALIGN_LEFT
                    elif col_idx == 2:
                        cell.alignment = ALIGN_CENTER
                    else:
                        cell.alignment = ALIGN_RIGHT
                row_idx += 1

        self._auto_fit_columns(ws)

    # =========================================================================
    # 6. SHEET: Distributions
    # =========================================================================
    def _build_distributions_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title="Distributions")
        ws.views.sheetView[0].showGridLines = True

        title_cell = ws.cell(row=1, column=1, value="Subject-wise Marks Range Distribution (% of Total Marks)")
        title_cell.font = SECTION_TITLE_FONT

        d_headers = [
            "Sr.No", "Subject", "Max Marks", "Absent(0)", "90-100", "80-89",
            "70-79", "60-69", "50-59", "40-49", "Below 40 (Fail)",
            "Total Students", "Highest %", "Lowest %", "Average %"
        ]

        for c_idx, h in enumerate(d_headers, 1):
            c = ws.cell(row=3, column=c_idx, value=h)
            c.fill = TABLE_HEADER_FILL
            c.font = TABLE_HEADER_FONT
            c.border = THIN_BORDER
            c.alignment = ALIGN_CENTER

        all_courses = self._get_theory_courses() + self._get_lab_courses()

        for sr_idx, c in enumerate(all_courses, 1):
            r_idx = 3 + sr_idx
            code = c["code"]
            name = c.get("name", code)

            enrolled = self._get_enrolled_students(code)
            if not enrolled:
                continue

            max_m = self._get_course_head_max(c, "TOT")

            marks_list = []
            pct_list = []
            abs_count = 0
            b90 = 0
            b80 = 0
            b70 = 0
            b60 = 0
            b50 = 0
            b40 = 0
            b_fail = 0

            for s in enrolled:
                if self._is_student_absent_in_course(s, c):
                    abs_count += 1
                    continue

                m_val = self._get_student_mark(s, code, "TOT")
                if m_val is None:
                    m_val = self._get_student_mark(s, code, "TH") or self._get_student_mark(s, code, "PR") or self._get_student_mark(s, code, "TW")

                is_failed = self._is_student_failed_in_course(s, c)

                if is_failed:
                    b_fail += 1
                    if m_val is not None and max_m > 0:
                        pct = (m_val / max_m) * 100.0
                        pct_list.append(pct)
                elif m_val is not None and max_m > 0:
                    pct = (m_val / max_m) * 100.0
                    pct_list.append(pct)
                    if pct >= 90:
                        b90 += 1
                    elif pct >= 80:
                        b80 += 1
                    elif pct >= 70:
                        b70 += 1
                    elif pct >= 60:
                        b60 += 1
                    elif pct >= 50:
                        b50 += 1
                    elif pct >= 40:
                        b40 += 1
                    else:
                        b_fail += 1
                else:
                    b_fail += 1

            total_enrolled = len(enrolled)
            highest_pct = round(max(pct_list), 2) if pct_list else 0.0
            lowest_pct = round(min(pct_list), 2) if pct_list else 0.0
            avg_pct = round(sum(pct_list) / len(pct_list), 2) if pct_list else 0.0

            vals = [
                sr_idx, name, max_m, abs_count, b90, b80, b70, b60, b50, b40, b_fail,
                total_enrolled, highest_pct, lowest_pct, avg_pct
            ]

            for col_idx, val in enumerate(vals, 1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if col_idx == 2:
                    cell.alignment = ALIGN_LEFT
                elif col_idx in [1, 3]:
                    cell.alignment = ALIGN_CENTER
                else:
                    cell.alignment = ALIGN_RIGHT

        self._auto_fit_columns(ws)

    # =========================================================================
    # Helpers
    # =========================================================================
    def _key_exists(self, key_name: str) -> bool:
        """Check if any student record contains key_name."""
        if not self.students:
            return False
        return key_name in self.students[0]

    def _auto_fit_columns(self, ws: openpyxl.worksheet.worksheet.Worksheet):
        """Auto-adjust column widths based on contents with safety bounds."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    lines = str(cell.value).split("\n")
                    for line in lines:
                        if len(line) > max_len:
                            max_len = len(line)
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 48), 12)
